"""Catalog loaders (SharePoint spec §8, PI-07 / PI-08): vendor price lists from a designated folder or an in-app
upload, the one-time bootstrap from the Pricing Intelligence dashboard's embedded catalog, and the hygiene pass that
marks duplicate / stale / superseded rows. Nothing here reads PTT, SL, SharePoint or the share for writing — price
lists are third-party data and the only spreadsheet input the spec allows (loaded by a loader, never hand-copied).

`refresh_all_step(run)` is picked up by `manage.py refresh_all`; it is a no-op until a vendor price folder is
configured (`PCA_VENDOR_PRICE_DIR` in `.env`, or `settings.ESTIMATING_VENDOR_PRICE_DIR`)."""

import csv
import io
import json
import os
from datetime import date, datetime
from decimal import Decimal
from pathlib import Path

from django.conf import settings
from django.db import connection, transaction
from django.utils import timezone

from apps.ingestion.bulk import bulk_update_from_values, fetch_dict, upsert

from . import rules
from .models import CatalogItem, CatalogSource, HygieneReport

TABLE = "estimating_catalogitem"
ITEM_COLS = ["source_id", "manufacturer", "manufacturer_raw", "part", "part_norm", "description", "cost", "msrp", "map_price",
             "category", "date_key", "confidence", "archived", "archive_reason", "superseded_by_id", "created_at"]
CONFLICT = ["source_id", "part_norm", "manufacturer_raw"]
FILE_EXTS = (".xlsx", ".xlsm", ".csv", ".tsv", ".txt")
PREVIEW_CAP = 500
MODES = ("upsert", "add-only", "prices-only")
PRICE_CAP = Decimal("1000000000")


def vendor_price_dir():
    p = os.environ.get("PCA_VENDOR_PRICE_DIR") or getattr(settings, "ESTIMATING_VENDOR_PRICE_DIR", "") or ""
    return Path(p).expanduser() if p else None


def refresh_all_step(run):
    """Called by refresh_all. Loads new / changed price files from the configured folder; otherwise reports why not."""
    folder = vendor_price_dir()
    if not folder or not folder.is_dir():
        return {"skipped": "no vendor price folder configured (PCA_VENDOR_PRICE_DIR)"}
    return load_folder(folder, run=run)


# ----------------------------------------------------------------------------------------------------------------
# helpers
# ----------------------------------------------------------------------------------------------------------------
def _price(v):
    """Catalog prices: 0 / blank / None all mean "unknown" (the dashboard's `r[4]||null`)."""
    if v in (None, "", 0, 0.0, False):
        return None
    try:
        d = Decimal(str(v))
    except Exception:  # noqa
        return None
    # a price of a billion dollars is a UPC / SKU that landed in a price column (the vendor files carry a few)
    return d if 0 < d < PRICE_CAP else None


def _today_key():
    t = timezone.localdate()
    return t.year * 10000 + t.month * 100 + t.day


def _cellstr(v):
    if v is None:
        return ""
    if isinstance(v, datetime):
        return v.strftime("%Y-%m-%d")
    if isinstance(v, date):
        return v.isoformat()
    if isinstance(v, float) and v.is_integer():
        return str(int(v))
    return str(v)


def _item_row(source_id, mfr_raw, part, norm, desc, cost, msrp, map_price, category, date_key, now):
    mfr = rules.clean_manufacturer(mfr_raw)
    conf = rules.confidence(cost, msrp, desc, rules.date_key_year(date_key), map_price)
    return (source_id, mfr[:120], (mfr_raw or "")[:200], part[:160], norm[:160], desc or "", cost, msrp, map_price,
            (category or "")[:255], int(date_key or 0), conf, False, "", None, now)


def _refresh_source_stats():
    """rows + dominant manufacturer per source, from the items."""
    with connection.cursor() as cur:
        cur.execute("""UPDATE estimating_catalogsource s SET rows = x.n, vendor = COALESCE(x.vendor, s.vendor)
                       FROM (SELECT source_id, COUNT(*) n, mode() WITHIN GROUP (ORDER BY manufacturer) vendor
                             FROM estimating_catalogitem GROUP BY source_id) x WHERE x.source_id = s.id""")
        cur.execute("UPDATE estimating_catalogsource SET rows = 0 WHERE id NOT IN (SELECT DISTINCT source_id FROM estimating_catalogitem)")


# ----------------------------------------------------------------------------------------------------------------
# hygiene
# ----------------------------------------------------------------------------------------------------------------
def run_hygiene(trigger="manual", notes=None, log=None):
    """The V100 current-only catalog as a loader rule: duplicate part numbers collapse to the better row, rows priced
    before 2023 are hidden, older explicit versions (V2 / MK II / GEN 3 / REV B) of a family are archived. Nothing is
    deleted — losers get `archived`, `archive_reason` and `superseded_by`. Returns the HygieneReport."""
    log = log or (lambda s: None)
    rows = []
    with connection.cursor() as cur:
        cur.execute("SELECT id, part, part_norm, manufacturer, date_key, cost, msrp, description <> '' FROM estimating_catalogitem")
        for i, p, n, m, dk, c, ms, d in cur.fetchall():
            rows.append({"id": i, "part": p, "part_norm": n, "manufacturer": m, "year": rules.date_key_year(dk), "date_key": dk,
                         "cost": c, "msrp": ms, "description": d})
    log("hygiene: %d rows" % len(rows))
    res = rules.hygiene(rows)
    updates = [(lid, True, "duplicate", wid) for lid, wid in res["duplicate"]]
    updates += [(i, True, "pre2023", None) for i in res["pre2023"]]
    updates += [(lid, True, "version", wid) for lid, wid in res["version"]]
    with transaction.atomic():
        with connection.cursor() as cur:
            cur.execute("UPDATE estimating_catalogitem SET archived = false, archive_reason = '', superseded_by_id = NULL WHERE archived")
        bulk_update_from_values(TABLE, "id", ["archived", "archive_reason", "superseded_by_id"], updates, page_size=5000)
    counts = dict(res["counts"])
    extra = fetch_dict("""SELECT COUNT(DISTINCT manufacturer) mfrs, COUNT(DISTINCT source_id) sources,
                                 COUNT(*) FILTER (WHERE manufacturer = 'Other') other_rows,
                                 COUNT(*) FILTER (WHERE manufacturer <> manufacturer_raw AND manufacturer <> 'Other') aliased_rows,
                                 COUNT(*) FILTER (WHERE cost IS NOT NULL) with_cost, COUNT(*) FILTER (WHERE msrp IS NOT NULL) with_msrp,
                                 COUNT(*) FILTER (WHERE map_price IS NOT NULL) with_map, COUNT(*) FILTER (WHERE date_key = 0) undated
                          FROM estimating_catalogitem WHERE NOT archived""")[0]
    counts.update({k: int(v or 0) for k, v in extra.items()})
    report = HygieneReport.objects.create(trigger=trigger, counts=counts, notes=notes or [])
    log("hygiene: %(current)s current · %(duplicates)s duplicates hidden · %(versions)s older versions archived · %(pre2023)s pre-2023 rows removed" % counts)
    return report


# ----------------------------------------------------------------------------------------------------------------
# the one-time bootstrap from the dashboard's embedded catalog
# ----------------------------------------------------------------------------------------------------------------
def bootstrap_from_dashboard(path, dry_run=False, log=None, batch=50000):
    """Seed the catalog from `const _D={...}` in pricing_intelligence.html (M manufacturers, S sources, C categories,
    R records [mfr_idx, part, norm, desc, cost, msrp, map, cat_idx, src_idx]). A sanctioned one-time seed of third-party
    vendor pricing (spec §8; AGENT_BRIEF "catalog bootstrap") — idempotent: re-running inserts nothing that exists."""
    log = log or (lambda s: None)
    path = Path(path)
    log("reading %s (%.1f MB)" % (path.name, path.stat().st_size / 1e6))
    D = None
    with open(path, "r", encoding="utf-8", errors="replace") as f:
        for line in f:
            if line.startswith("const _D="):
                text = line[len("const _D="):].strip()
                if text.endswith(";"):
                    text = text[:-1]
                D = json.loads(text)
                break
    if D is None:
        raise ValueError("no `const _D=` line found in %s" % path)
    M, S, C, R = D["M"], D["S"], D["C"], D["R"]
    log("%d manufacturers, %d sources, %d categories, %d records" % (len(M), len(S), len(C), len(R)))
    src_dates = [rules.source_date(s) for s in S]
    if dry_run:
        return {"manufacturers": len(M), "sources": len(S), "categories": len(C), "records": len(R),
                "dated_sources": sum(1 for d in src_dates if d["key"]), "sample": R[:3]}
    now = timezone.now()
    src_ids = []
    with transaction.atomic():
        for name, info in zip(S, src_dates):
            src, _ = CatalogSource.objects.get_or_create(
                name=name[:255], defaults={"kind": CatalogSource.Kind.BOOTSTRAP, "path": str(path), "date_key": info["key"],
                                           "date_label": info["label"], "date_precision": info["precision"]})
            src_ids.append(src.id)
    inserted = blank = 0
    rows = []
    for n, r in enumerate(R):
        mi, part, norm, desc, cost, msrp, mp, ci, si = (r + [None] * 9)[:9]
        part = str(part or "")
        norm = str(norm or "") or rules.norm_part(part)
        if not norm or si is None or not (0 <= si < len(src_ids)):
            blank += 1
            continue
        mfr_raw = M[mi] if mi is not None and 0 <= mi < len(M) else ""
        cat = C[ci] if ci is not None and 0 <= ci < len(C) else ""
        rows.append(_item_row(src_ids[si], mfr_raw, part, norm, str(desc or ""), _price(cost), _price(msrp), _price(mp), cat, src_dates[si]["key"], now))
        if len(rows) >= batch:
            inserted += upsert(TABLE, ITEM_COLS, rows, CONFLICT)
            rows = []
            log("  %d / %d records → %d inserted" % (n + 1, len(R), inserted))
    if rows:
        inserted += upsert(TABLE, ITEM_COLS, rows, CONFLICT)
    log("inserted %d rows (%d skipped: blank part / bad source index; %d collapsed inside a file on the same part + manufacturer)"
        % (inserted, blank, len(R) - blank - inserted))
    _refresh_source_stats()
    report = run_hygiene(trigger="bootstrap", log=log,
                         notes=["bootstrap from %s: %d records, %d inserted" % (path.name, len(R), inserted)])
    return {"manufacturers": len(M), "sources": len(S), "categories": len(C), "records": len(R), "inserted": inserted,
            "skipped": blank, "hygiene": report.counts}


# ----------------------------------------------------------------------------------------------------------------
# files: parse, apply, folder walk, removal
# ----------------------------------------------------------------------------------------------------------------
def parse_file(path=None, data=None, name=""):
    """A vendor price list / BOM workbook → (rows, sheets_found, errors, sheets). rows come from rules.catalog_rows
    (header detection with weighted aliases); sheets = [(title, matrix)] for the estimate importer."""
    name = name or (Path(path).name if path else "upload")
    ext = name.lower().rsplit(".", 1)[-1] if "." in name else ""
    sheets, errors = [], []
    if ext in ("xlsx", "xlsm"):
        from openpyxl import load_workbook
        try:
            wb = load_workbook(filename=(path or io.BytesIO(data)), read_only=True, data_only=True)
            for ws in wb.worksheets:
                sheets.append((ws.title, [[_cellstr(v) for v in row] for row in ws.iter_rows(values_only=True)]))
            wb.close()
        except Exception as e:  # noqa
            errors.append("%s: %s" % (name, e))
    elif ext in ("csv", "tsv", "txt"):
        raw = data if data is not None else Path(path).read_bytes()
        text = raw.decode("utf-8-sig", errors="replace")
        delim = "\t" if ext == "tsv" else ","
        try:
            delim = csv.Sniffer().sniff(text[:5000], delimiters=",;\t|").delimiter
        except Exception:  # noqa
            pass
        sheets.append((name, list(csv.reader(io.StringIO(text), delimiter=delim))))
    elif ext == "xls":
        errors.append("%s: legacy .xls workbooks are not supported — save it as .xlsx" % name)
    else:
        errors.append("%s: unsupported file type" % name)
    rows, found = [], 0
    for title, matrix in sheets:
        r = rules.catalog_rows(matrix, name, title)
        if r:
            found += 1
            rows.extend(r)
    return rows, found, errors, sheets


def preview_rows(rows):
    """Import stats the way the dashboard showed them: rows after dedupe, brands, existing matches, new products."""
    rows = rules.dedupe_catalog_rows(rows)
    norms = list({rules.norm_part(r["part"]) for r in rows if rules.norm_part(r["part"])})
    existing = set()
    if norms:
        existing = {r["part_norm"] for r in fetch_dict("SELECT DISTINCT part_norm FROM estimating_catalogitem WHERE NOT archived AND part_norm = ANY(%s)", [norms])}
    return rows, {"rows": len(rows), "brands": len({r["mfr"] for r in rows}), "existing": len(existing & set(norms)),
                  "new": len(set(norms) - existing), "stale": sum(1 for r in rows if rules.is_pre_2023(rules.date_key_year(rules.source_date(r.get("source_date") or r.get("file") or "")["key"])))}


def apply_catalog_rows(rows, name, mode="upsert", dry_run=True, kind="upload", path="", mtime=None, size=None, account=None,
                       hygiene=True, log=None):
    """Apply parsed rows as a CatalogSource. Matching is by normalized part number only (as the dashboard did).
    Modes: upsert (add new parts, re-price existing ones), add-only (new parts only), prices-only (existing only).
    Rows dated before 2023 are skipped. A re-priced part is a NEW row under this source dated today when the file
    carries no date — the hygiene pass then makes it the current row and marks the old one superseded, so the price
    history survives (the dashboard overwrote in place). dry_run returns the counts and a preview, writes nothing."""
    log = log or (lambda s: None)
    if mode not in MODES:
        mode = "upsert"
    rows = rules.dedupe_catalog_rows(rows)
    products = [rules.imported_product(r) for r in rows]
    for p, r in zip(products, rows):
        p["qty"] = r.get("qty", 1)
        p["area"] = r.get("area", "")
        p["labor"] = r.get("labor") or {}
    norms = list({p["part_norm"] for p in products if p["part_norm"]})
    existing = {}
    if norms:
        for r in fetch_dict("""SELECT id, part_norm, manufacturer, manufacturer_raw, description, cost, msrp, map_price, date_key
                               FROM estimating_catalogitem WHERE NOT archived AND part_norm = ANY(%s)""", [norms]):
            existing[r["part_norm"]] = r
    file_info = rules.source_date(name)
    fallback_key = file_info["key"] or _today_key()
    counts = {"rows": len(rows), "added": 0, "updated": 0, "skipped": 0, "stale": 0, "unchanged": 0, "blank": 0,
              "existing": len(existing), "new": len(set(norms) - set(existing)), "mode": mode}
    to_insert, preview = [], []
    for p in products:
        if not p["part_norm"]:
            counts["blank"] += 1
            counts["skipped"] += 1
            continue
        if not p["date_key"]:
            p["date_key"] = fallback_key
        if rules.is_pre_2023(rules.date_key_year(p["date_key"])):
            counts["stale"] += 1
            counts["skipped"] += 1
            continue
        old = existing.get(p["part_norm"])
        if old:
            if mode == "add-only":
                counts["skipped"] += 1
                continue
            changed = False
            if p["cost"] is not None and p["cost"] != old["cost"]:
                changed = True
            if p["msrp"] is not None and p["msrp"] != old["msrp"]:
                changed = True
            if p["map"] is not None and p["map"] != old["map_price"]:
                changed = True
            if p["description"] and p["description"] != p["part"] and p["description"] != (old["description"] or ""):
                changed = True
            if (old["manufacturer"] in ("", "Other", "Imported")) and p["manufacturer"] not in ("", "Other", "Imported"):
                changed = True
            if not changed:
                counts["unchanged"] += 1
                counts["skipped"] += 1
                continue
            # carry what the file did not say so the new row can be the current one without losing data
            for k, ok in (("cost", "cost"), ("msrp", "msrp"), ("map", "map_price")):
                if p[k] is None and old[ok] is not None:
                    p[k] = old[ok]
            if (not p["description"] or p["description"] == p["part"]) and old["description"]:
                p["description"] = old["description"]
            if p["manufacturer"] in ("", "Other", "Imported") and old["manufacturer"] not in ("", "Other", "Imported"):
                p["manufacturer"], p["manufacturer_raw"] = old["manufacturer"], old["manufacturer_raw"]
            p["action"] = "update"
            counts["updated"] += 1
        else:
            if mode == "prices-only":
                counts["skipped"] += 1
                continue
            p["action"] = "add"
            counts["added"] += 1
        to_insert.append(p)
        if len(preview) < PREVIEW_CAP:
            preview.append({k: (str(v) if isinstance(v, Decimal) else v) for k, v in p.items() if k != "labor"})
    if dry_run:
        return {"counts": counts, "preview": preview, "mode": mode, "products": to_insert}
    now = timezone.now()
    with transaction.atomic():
        src, created = CatalogSource.objects.get_or_create(
            name=name[:255], defaults={"kind": kind, "path": path, "date_key": file_info["key"], "date_label": file_info["label"],
                                       "date_precision": file_info["precision"], "mode": mode})
        if not created:
            CatalogItem.objects.filter(source=src).delete()        # a re-loaded file replaces its own rows
        src.kind, src.path, src.mode, src.file_mtime, src.file_size, src.loaded_by = kind, path or src.path, mode, mtime, size, account
        item_rows = [_item_row(src.id, p["manufacturer_raw"], p["part"], p["part_norm"], p["description"], p["cost"], p["msrp"],
                               p["map"], p["category"], p["date_key"], now) for p in to_insert]
        n = upsert(TABLE, ITEM_COLS, item_rows, CONFLICT)
        counts["inserted"] = n
        src.rows = n
        src.counts = counts
        src.save()
    _refresh_source_stats()
    out = {"counts": counts, "source_id": src.id, "mode": mode}
    if hygiene:
        report = run_hygiene(trigger=kind, log=log, notes=["%s (%s): +%d new, %d re-priced, %d skipped" % (name, mode, counts["added"], counts["updated"], counts["skipped"])])
        out["hygiene"] = report.counts
    log("%s: %d rows → %d added, %d re-priced, %d skipped (%d stale, %d unchanged)" % (name, len(rows), counts["added"], counts["updated"], counts["skipped"], counts["stale"], counts["unchanged"]))
    return out


def load_folder(folder, run=None, mode="upsert", force=False, log=None):
    """Load every new or changed price file (by name + mtime + size) from the vendor price folder, one hygiene pass at the end."""
    from apps.ingestion.loaders import issue
    log = log or (lambda s: None)
    folder = Path(folder)
    results, errors = [], []
    for f in sorted(folder.iterdir()):
        if not f.is_file() or f.name.startswith("~$") or f.suffix.lower() not in FILE_EXTS:
            continue
        st = f.stat()
        src = CatalogSource.objects.filter(name=f.name[:255]).first()
        if src and not force and src.file_mtime == st.st_mtime and src.file_size == st.st_size:
            continue
        rows, found, errs, _ = parse_file(path=f)
        errors.extend(errs)
        if not rows:
            errors.append("%s: no price rows recognised (no Part / Description header found)" % f.name)
            continue
        res = apply_catalog_rows(rows, f.name, mode=mode, dry_run=False, kind="file", path=str(f), mtime=st.st_mtime, size=st.st_size,
                                 hygiene=False, log=log)
        results.append({"file": f.name, **res["counts"]})
    if results:
        run_hygiene(trigger="folder", log=log, notes=["%s: %d file(s)" % (folder, len(results))])
    if run is not None:
        for e in errors:
            issue(run, "estimating_price_file", "warning", source_system="share", source_key=e[:120], message=e)
    return {"folder": str(folder), "files": len(results), "results": results, "errors": errors}


def remove_source(source_id, log=None):
    """Undo an upload / file load: drop its rows and let the hygiene pass restore the rows it superseded."""
    src = CatalogSource.objects.filter(pk=source_id).first()
    if src is None:
        return None
    name, n = src.name, src.items.count()
    src.delete()
    _refresh_source_stats()
    run_hygiene(trigger="remove", log=log, notes=["removed %s (%d rows)" % (name, n)])
    return {"name": name, "rows": n}
