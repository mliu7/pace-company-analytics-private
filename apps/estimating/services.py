"""Estimate services (SharePoint spec §8): the rate card in force, line / room / grand maths through rules.py, the
versioned save with a stale-form check, duplicate, quick-add, workbook import with catalog price lookup, and the
exports (multi-room xlsx with a Summary sheet, detail CSV, clipboard text). Views call these; nothing here renders."""

import csv
import io
from datetime import datetime
from decimal import Decimal

from django.db import transaction
from django.utils import timezone

from apps.ingestion.bulk import fetch_dict

from . import rules
from .models import CatalogItem, Estimate, EstimateVersion, LaborRate, Line, Room

D0 = Decimal("0")


class Stale(Exception):
    """The form was opened before someone else saved (spec §2: "changed by X at hh:mm — reload")."""

    def __init__(self, est):
        who = est.updated_by.display_name if est.updated_by_id and est.updated_by else "someone"
        when = timezone.localtime(est.updated_at).strftime("%-I:%M %p") if est.updated_at else ""
        super().__init__("changed by %s at %s — reload before saving" % (who, when))
        self.who, self.when = who, when


# ----------------------------------------------------------------------------------------------------------------
# rate card
# ----------------------------------------------------------------------------------------------------------------
def current_rates(on=None):
    """{rate id: {label, group, cost, sell, effective_from, order}} — the latest row effective on `on` (today);
    defaults fill anything the table lacks."""
    on = on or timezone.localdate()
    out = {rid: {"rate_id": rid, "label": v["label"], "group": v["group"], "cost": v["cost"], "sell": v["sell"], "effective_from": None, "order": i, "default": True}
           for i, (rid, v) in enumerate(rules.DEFAULT_RATE_MAP.items())}
    seen = set()
    for r in LaborRate.objects.filter(effective_from__lte=on).order_by("rate_id", "-effective_from"):
        if r.rate_id in seen:
            continue
        seen.add(r.rate_id)
        d = rules.DEFAULT_RATE_MAP.get(r.rate_id, {})
        out[r.rate_id] = {"rate_id": r.rate_id, "label": r.label, "group": r.group, "cost": r.cost, "sell": r.sell, "effective_from": r.effective_from,
                          "order": r.order, "default": bool(d and d["cost"] == r.cost and d["sell"] == r.sell)}
    return out


def rates_list(rates):
    rows = sorted(rates.values(), key=lambda r: r["order"])
    for r in rows:
        r["margin"] = rules.rate_margin(r["cost"], r["sell"])
        r["markup"] = float(rules.q3(r["sell"] / r["cost"])) if r["cost"] else None
        r["group_label"] = rules.GROUP_LABELS.get(r["group"], r["group"])
        r["default_cost"] = rules.DEFAULT_RATE_MAP.get(r["rate_id"], {}).get("cost")
        r["default_sell"] = rules.DEFAULT_RATE_MAP.get(r["rate_id"], {}).get("sell")
    return rows


def rates_json(rates):
    return {rid: {"label": r["label"], "group": r["group"], "cost": float(r["cost"]), "sell": float(r["sell"])} for rid, r in rates.items()}


def rates_history(limit=200):
    return list(LaborRate.objects.select_related("set_by").order_by("-effective_from", "order")[:limit])


# ----------------------------------------------------------------------------------------------------------------
# payload ↔ rows
# ----------------------------------------------------------------------------------------------------------------
def jsonable(v):
    if isinstance(v, Decimal):
        return float(v)
    if isinstance(v, (datetime,)):
        return v.isoformat()
    if isinstance(v, dict):
        return {k: jsonable(x) for k, x in v.items()}
    if isinstance(v, (list, tuple)):
        return [jsonable(x) for x in v]
    if hasattr(v, "isoformat"):
        return v.isoformat()
    return v


def item_snapshot(it):
    """The copied item fields a line keeps from a CatalogItem."""
    return {"catalog_item_id": it.id, "manufacturer": it.manufacturer, "part": it.part, "part_norm": it.part_norm, "description": it.description,
            "category": it.category, "source_name": it.source.name if it.source_id else "", "date_key": it.date_key,
            "item_msrp": it.msrp, "item_map": it.map_price, "cost": it.cost or D0}


def normalize_line(d, title="", prev_area=""):
    """One builder / import line dict → the shape the maths and the Line model expect. Sell is authoritative; markup
    follows sell (or sell follows markup when only markup came)."""
    kind = "note" if (d.get("kind") == "note") else "item"
    cost = rules.dec(d.get("cost"))
    cost = cost if cost > 0 else D0
    msrp = rules.dec(d.get("item_msrp")) if d.get("item_msrp") not in (None, "") else None
    mk_in = d.get("markup")
    sell_in = d.get("sell")
    if sell_in not in (None, ""):
        sell = rules.q2(sell_in)
        markup = rules.markup_from_sell(cost, sell) if cost > 0 else (rules.q3(mk_in) if mk_in not in (None, "") else None)
    elif mk_in not in (None, ""):
        markup = rules.q3(mk_in)
        sell = rules.sell_from_markup(cost, markup) if cost > 0 else rules.q2(msrp or 0)
    else:
        sell = rules.default_sell(cost, msrp)
        markup = rules.MATERIAL_MARKUP_DEFAULT if cost > 0 else None
    if kind == "note":
        cost, sell, markup = D0, D0, None
    hours_in = d.get("hours") or {}
    hours = {}
    for t in rules.LABOR_TYPES:
        h = rules.dec(hours_in.get(t["id"], d.get("lh_" + t["id"])))
        hours[t["id"]] = h.quantize(rules.HOURS_Q) if h > 0 else D0
    part = rules.clean(d.get("part"))[:160]
    misc = bool(d["misc"]) if d.get("misc") not in (None, "") else rules.misc_like(d.get("category"), d.get("manufacturer"), part, d.get("description"))
    return {"id": d.get("id"), "catalog_item_id": d.get("catalog_item_id") or None, "kind": kind,
            "area": rules.clean(d.get("area"))[:120] or prev_area or rules.clean(title)[:120] or "PACE Estimate",
            "manufacturer": rules.clean(d.get("manufacturer"))[:120], "part": part, "part_norm": rules.norm_part(part)[:160],
            "description": rules.clean(d.get("description")), "category": rules.clean(d.get("category"))[:255],
            "source_name": rules.clean(d.get("source_name"))[:255], "date_key": int(rules.dec(d.get("date_key"))),
            "item_msrp": msrp, "item_map": rules.dec(d.get("item_map")) if d.get("item_map") not in (None, "") else None,
            "qty": rules.line_qty(d.get("qty")), "cost": cost, "markup": markup, "sell": sell, "misc": misc, "hours": hours}


def compute(rooms, rates):
    """rooms: [{name, order, notes, lines:[normalized line]}] → adds line totals, room totals, grand totals, warnings."""
    for room in rooms:
        for ln in room["lines"]:
            ln["totals"] = rules.line_totals(ln["cost"], ln["qty"], ln["sell"], ln["hours"], rates, markup=ln["markup"])
        room["totals"] = rules.sum_totals([ln["totals"] for ln in room["lines"]])
    grand = rules.sum_totals([r["totals"] for r in rooms])
    flat = [{"kind": ln["kind"], "misc": ln["misc"], "cost": ln["cost"], "markup": ln["markup"], "qty": ln["qty"], "hours": ln["hours"],
             "part": ln["part"], "description": ln["description"]} for r in rooms for ln in r["lines"]]
    return {"rooms": rooms, "grand": grand, "warnings": rules.policy_warnings(grand, flat)}


def rooms_from_payload(payload, title=""):
    rooms = []
    for i, r in enumerate(payload.get("rooms") or []):
        name = rules.clean(r.get("name"))[:120] or ("Room %d" % (i + 1))
        lines, prev = [], ""
        for j, d in enumerate(r.get("lines") or []):
            ln = normalize_line(d, title=title, prev_area=prev or name)
            ln["order"] = j
            prev = ln["area"]
            lines.append(ln)
        rooms.append({"id": r.get("id"), "name": name, "order": i, "notes": rules.clean(r.get("notes")), "lines": lines})
    return rooms


def rooms_from_db(est):
    rooms = []
    for room in est.rooms.all().prefetch_related("lines"):
        lines = []
        for ln in room.lines.all():
            lines.append({"id": ln.id, "catalog_item_id": ln.catalog_item_id, "kind": ln.kind, "order": ln.order, "area": ln.area,
                          "manufacturer": ln.manufacturer, "part": ln.part, "part_norm": ln.part_norm, "description": ln.description,
                          "category": ln.category, "source_name": ln.source_name, "date_key": ln.date_key, "item_msrp": ln.item_msrp,
                          "item_map": ln.item_map, "qty": ln.qty, "cost": ln.cost, "markup": ln.markup, "sell": ln.sell, "misc": ln.misc,
                          "hours": ln.hours()})
        rooms.append({"id": room.id, "name": room.name, "order": room.order, "notes": room.notes, "lines": lines})
    return rooms


def estimate_payload(est, rates=None):
    """Everything the builder page needs, computed live against the rate card in force."""
    rates = rates or current_rates()
    comp = compute(rooms_from_db(est), rates)
    for room in comp["rooms"]:
        for ln in room["lines"]:
            ln["date_label"] = rules.date_key_label(ln["date_key"])
    bid = None
    if est.bid_id and est.bid:
        b = est.bid
        bid = {"id": b.id, "project_name": b.project_name, "client_name": b.client_name, "stage": b.stage, "budget": b.budget, "value": b.value,
               "job_number_raw": b.job_number_raw, "bid_due": b.bid_due, "estimator": b.estimator.canonical_name if b.estimator_id and b.estimator else ""}
    return jsonable({
        "id": est.id, "title": est.title, "client_name": est.client_name, "customer_id": est.customer_id,
        "customer_name": est.customer.canonical_name if est.customer_id and est.customer else "", "notes": est.notes,
        "bid_id": est.bid_id, "bid": bid, "project_id": est.project_id,
        "project_number": est.project.display_number if est.project_id and est.project else "", "status": est.status,
        "version_no": est.version_no, "owner": est.owner.display_name if est.owner_id and est.owner else "",
        "updated_at": est.updated_at, "updated_by": est.updated_by.display_name if est.updated_by_id and est.updated_by else "",
        "approval_ref": est.approval_ref, "source_file": est.source_file,
        "rooms": comp["rooms"], "grand": comp["grand"], "warnings": comp["warnings"], "rates": rates_json(rates),
        "labor_types": rules.LABOR_TYPES, "markup_default": rules.MATERIAL_MARKUP_DEFAULT,
    })


def _apply_meta(est, payload):
    from apps.core.models import Customer, Project
    if "title" in payload:
        est.title = rules.clean(payload.get("title"))[:200]
    if "client_name" in payload:
        est.client_name = rules.clean(payload.get("client_name"))[:200]
    if "notes" in payload:
        est.notes = str(payload.get("notes") or "")
    if "status" in payload and payload["status"] in dict(Estimate.Status.choices):
        est.status = payload["status"]
    if "customer_id" in payload:
        cid = payload.get("customer_id")
        est.customer = Customer.objects.filter(pk=cid).first() if cid else None
    if "bid_id" in payload:
        from apps.bids.models import Bid
        bid = payload.get("bid_id")
        est.bid = Bid.objects.filter(pk=bid).select_related("client", "project").first() if bid else None
        if est.bid and not est.client_name:
            est.client_name = est.bid.client_name[:200]
        if est.bid and est.bid.client_id and not est.customer_id:
            est.customer = est.bid.client
        if est.bid and est.bid.project_id and not est.project_id:
            est.project = est.bid.project
    if "project_number" in payload:
        pn = rules.clean(payload.get("project_number")).upper()
        est.project = Project.objects.filter(canonical_project_number=pn).first() if pn else None


def _write_rooms(est, comp):
    est.rooms.all().delete()
    for room in comp["rooms"]:
        r = Room.objects.create(estimate=est, name=room["name"], order=room["order"], notes=room["notes"])
        room["id"] = r.id
        objs = []
        for ln in room["lines"]:
            t = ln["totals"]
            kw = {rules.LABOR_COLUMN[k]: v for k, v in ln["hours"].items()}
            objs.append(Line(room=r, catalog_item_id=ln["catalog_item_id"], kind=ln["kind"], order=ln["order"], area=ln["area"],
                             manufacturer=ln["manufacturer"], part=ln["part"], part_norm=ln["part_norm"], description=ln["description"],
                             category=ln["category"], source_name=ln["source_name"], date_key=ln["date_key"], item_msrp=ln["item_msrp"],
                             item_map=ln["item_map"], qty=ln["qty"], cost=ln["cost"], markup=ln["markup"], sell=ln["sell"], misc=ln["misc"],
                             total_cost=rules.q2(t["total_cost"]), total_sell=rules.q2(t["total_sell"]), **kw))
        created = Line.objects.bulk_create(objs)
        for ln, obj in zip(room["lines"], created):
            ln["id"] = obj.id


def _cache_totals(est, comp):
    g = comp["grand"]
    est.total_cost, est.total_sell = rules.q2(g["total_cost"]), rules.q2(g["total_sell"])
    est.equipment_cost, est.equipment_sell = rules.q2(g["equipment_cost_ext"]), rules.q2(g["equipment_sell_ext"])
    est.labor_cost, est.labor_sell = rules.q2(g["labor_cost"]), rules.q2(g["labor_sell"])
    est.labor_hours = rules.dec(g["labor_hours"]).quantize(rules.HOURS_Q)
    est.line_count, est.room_count = int(g["lines"]), len(comp["rooms"])


def _snapshot(est, comp, rates):
    return jsonable({"meta": {"title": est.title, "client_name": est.client_name, "notes": est.notes, "status": est.status, "bid_id": est.bid_id,
                              "customer_id": est.customer_id, "project_id": est.project_id},
                     "rooms": [{"name": r["name"], "order": r["order"], "notes": r["notes"],
                                "lines": [{k: v for k, v in ln.items() if k != "totals"} for ln in r["lines"]],
                                "totals": {k: v for k, v in r["totals"].items() if k != "hours_by_type"}} for r in comp["rooms"]],
                     "rates": rates_json(rates), "totals": {k: v for k, v in comp["grand"].items() if k != "hours_by_type"},
                     "warnings": comp["warnings"]})


def _version(est, comp, rates, account, note=""):
    est.version_no += 1
    est.updated_by = account
    est.save()
    EstimateVersion.objects.create(estimate=est, version_no=est.version_no, snapshot=_snapshot(est, comp, rates),
                                   totals=jsonable({k: v for k, v in comp["grand"].items() if k != "hours_by_type"}), saved_by=account, note=note[:200])


def save_estimate(est, payload, account, note="Saved"):
    """Replace the estimate's rooms / lines from the builder payload, recompute at today's rate card, cache totals,
    bump the version and store a snapshot. Refuses a stale form (payload.version_no ≠ current) with Stale."""
    rates = current_rates()
    with transaction.atomic():
        est = Estimate.objects.select_for_update().get(pk=est.pk)   # no select_related: a lock cannot span the nullable join
        v = payload.get("version_no")
        if v not in (None, "") and int(v) != est.version_no:
            raise Stale(est)
        _apply_meta(est, payload)
        comp = compute(rooms_from_payload(payload, est.title), rates)
        _write_rooms(est, comp)
        _cache_totals(est, comp)
        _version(est, comp, rates, account, note)
    return est


def save_meta(est, payload, account):
    """Title / client / notes / bid / customer / project / status only — lines untouched (the builder's meta auto-save)."""
    rates = current_rates()
    with transaction.atomic():
        est = Estimate.objects.select_for_update().get(pk=est.pk)
        _apply_meta(est, payload)
        comp = compute(rooms_from_db(est), rates)
        _cache_totals(est, comp)
        _version(est, comp, rates, account, "Details changed")
    return est


def recompute(est, account, note):
    rates = current_rates()
    with transaction.atomic():
        est = Estimate.objects.select_for_update().get(pk=est.pk)
        comp = compute(rooms_from_db(est), rates)
        for room in comp["rooms"]:
            for ln in room["lines"]:
                Line.objects.filter(pk=ln["id"]).update(total_cost=rules.q2(ln["totals"]["total_cost"]), total_sell=rules.q2(ln["totals"]["total_sell"]))
        _cache_totals(est, comp)
        _version(est, comp, rates, account, note)
    return est


def new_estimate(account, title="", client_name="", bid_id=None):
    est = Estimate.objects.create(title=rules.clean(title)[:200], client_name=rules.clean(client_name)[:200], owner=account, updated_by=account)
    if bid_id:
        _apply_meta(est, {"bid_id": bid_id})
        if not est.title and est.bid:
            est.title = est.bid.project_name[:200]
    Room.objects.create(estimate=est, name="Room 1", order=0)
    est.room_count = 1
    est.save()
    return est


def duplicate_estimate(est, account):
    rates = current_rates()
    with transaction.atomic():
        copy = Estimate.objects.create(title=(est.title or str(est))[:195] + " Copy", client_name=est.client_name, customer=est.customer, notes=est.notes,
                                       bid=est.bid, project=est.project, owner=account, updated_by=account, status=Estimate.Status.DRAFT)
        comp = compute(rooms_from_db(est), rates)
        _write_rooms(copy, comp)
        _cache_totals(copy, comp)
        _version(copy, comp, rates, account, "Duplicated from #%d" % est.id)
    return copy


def add_line(est, account, item=None, room_id=None, qty=1, area="", part=""):
    """Quick add (search page, compare, quick-add box): one new line in the chosen (or first) room, sell = cost × 1.265."""
    rates = current_rates()
    with transaction.atomic():
        est = Estimate.objects.select_for_update().get(pk=est.pk)
        room = est.rooms.filter(pk=room_id).first() if room_id else None
        room = room or est.rooms.order_by("order", "id").first() or Room.objects.create(estimate=est, name="Room 1", order=0)
        last = room.lines.order_by("-order", "-id").first()
        d = item_snapshot(item) if item is not None else {"part": part, "description": part, "manufacturer": "Imported"}
        d["area"] = area or (last.area if last else "")
        d["qty"] = qty
        ln = normalize_line(d, title=est.title, prev_area=(last.area if last else room.name))
        ln["order"] = (last.order + 1) if last else 0
        ln["totals"] = rules.line_totals(ln["cost"], ln["qty"], ln["sell"], ln["hours"], rates, markup=ln["markup"])
        kw = {rules.LABOR_COLUMN[k]: v for k, v in ln["hours"].items()}
        line = Line.objects.create(room=room, catalog_item_id=ln["catalog_item_id"], kind=ln["kind"], order=ln["order"], area=ln["area"],
                                   manufacturer=ln["manufacturer"], part=ln["part"], part_norm=ln["part_norm"], description=ln["description"],
                                   category=ln["category"], source_name=ln["source_name"], date_key=ln["date_key"], item_msrp=ln["item_msrp"],
                                   item_map=ln["item_map"], qty=ln["qty"], cost=ln["cost"], markup=ln["markup"], sell=ln["sell"], misc=ln["misc"],
                                   total_cost=rules.q2(ln["totals"]["total_cost"]), total_sell=rules.q2(ln["totals"]["total_sell"]), **kw)
        comp = compute(rooms_from_db(est), rates)
        _cache_totals(est, comp)
        _version(est, comp, rates, account, "Added %s" % (ln["part"] or ln["description"])[:60])
    return line


# ----------------------------------------------------------------------------------------------------------------
# workbook import (PI-07 estimate import) with catalog price lookup by normalized part
# ----------------------------------------------------------------------------------------------------------------
def import_rooms(sheets, source_label):
    """sheets = [(title, matrix)] → {"rooms": [{name, lines}], "matched", "unmatched", "notes", "strategy"}. Strategy 1: any
    sheet with a Room / Area column groups by it; strategy 2: every sheet except values / summary is a room. Each line is
    price-matched by normalized part (+ manufacturer either-way substring when given); the file's cost overrides."""
    rooms = rules.estimate_rows_by_room_column(sheets)
    strategy = "room_column"
    if not rooms:
        strategy = "sheet_per_room"
        for title, matrix in sheets:
            if rules.clean(title).lower() in ("values", "summary"):
                continue
            r = rules.estimate_rows_from_sheet(title, matrix)
            if r:
                rooms.append(r)
    norms = list({rules.norm_part(ln["model"]) for r in rooms for ln in r["lines"] if not ln["note"] and rules.norm_part(ln["model"])})
    cands = {}
    if norms:
        rows = fetch_dict("""SELECT i.id, i.manufacturer, i.part, i.part_norm, i.description, i.cost, i.msrp, i.map_price, i.category, i.date_key,
                                    i.archived, s.name source_name
                             FROM estimating_catalogitem i JOIN estimating_catalogsource s ON s.id = i.source_id
                             WHERE i.part_norm = ANY(%s) ORDER BY i.archived, i.date_key DESC, i.id""", [norms])
        for r in rows:
            cands.setdefault(r["part_norm"], []).append(r)
    matched = unmatched = notes = 0
    out = []
    for i, r in enumerate(rooms):
        lines = []
        for j, ln in enumerate(r["lines"]):
            if ln["note"]:
                notes += 1
                lines.append({"kind": "note", "area": ln["room"], "manufacturer": "NOTE", "part": "", "description": ln["desc"], "qty": 1,
                              "cost": 0, "sell": 0, "category": "Note", "source_name": source_label, "hours": {}, "order": j})
                continue
            pn = rules.norm_part(ln["model"])
            mf = ln["mfr"].lower()
            match = None
            for c in cands.get(pn, []):
                cm = (c["manufacturer"] or "").lower()
                if not mf or (mf in cm or cm in mf):
                    match = c
                    break
            if match is None and cands.get(pn):
                match = cands[pn][0]
            if match:
                matched += 1
            else:
                unmatched += 1
            cost = ln["cost"] if ln["cost"] > 0 else rules.dec(match["cost"] if match else 0)
            markup = ln["markup"] if ln["markup"] > 0 else rules.MATERIAL_MARKUP_DEFAULT
            sell = ln["sell"] if ln["sell"] > 0 else (rules.sell_from_markup(cost, markup) if cost > 0 else rules.q2((match or {}).get("msrp") or 0))
            lines.append({"kind": "item", "catalog_item_id": match["id"] if match else None, "area": ln["room"],
                          "manufacturer": ln["mfr"] or (match["manufacturer"] if match else "Imported"), "part": ln["model"] or (match["part"] if match else ""),
                          "description": ln["desc"] or (match["description"] if match else ""), "category": ln["item"] or (match["category"] if match else ""),
                          "source_name": (match["source_name"] if match else source_label), "date_key": match["date_key"] if match else 0,
                          "item_msrp": match["msrp"] if match else None, "item_map": match["map_price"] if match else None,
                          "qty": max(1, int(ln["qty"] or 1)), "cost": cost, "markup": markup, "sell": sell, "hours": ln["labor"], "order": j,
                          "matched": bool(match)})
        out.append({"name": r["name"], "order": i, "notes": "", "lines": lines})
    return {"rooms": out, "matched": matched, "unmatched": unmatched, "notes": notes, "strategy": strategy}


# ----------------------------------------------------------------------------------------------------------------
# exports (PI-09): multi-room xlsx with Summary, detail CSV, clipboard text
# ----------------------------------------------------------------------------------------------------------------
def _money(v):
    return float(rules.q2(v))


def export_xlsx(est, payload):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter
    wb = Workbook()
    ws = wb.active
    ws.title = "Summary"
    bold = Font(bold=True)
    head_fill = PatternFill("solid", fgColor="FFF2CC")
    ws.append(["Project", est.title])
    ws.append(["Client", est.client_name])
    ws.append(["Notes", est.notes])
    ws.append(["Status", est.get_status_display()])
    ws.append(["Exported", timezone.localtime().strftime("%Y-%m-%d %H:%M")])
    ws.append([])
    hdr = ["Room", "Items", "Equipment Cost", "Equipment Sell", "Labor Hours", "Labor Cost", "Labor Sell", "Total Cost", "Total Sell", "Profit", "Margin"]
    ws.append(hdr)
    for c in ws[ws.max_row]:
        c.font, c.fill = bold, head_fill
    for r in payload["rooms"]:
        t = r["totals"]
        ws.append([r["name"], t["lines"], _money(t["equipment_cost_ext"]), _money(t["equipment_sell_ext"]), float(t["labor_hours"]), _money(t["labor_cost"]),
                   _money(t["labor_sell"]), _money(t["total_cost"]), _money(t["total_sell"]), _money(t["profit"]), (t["margin"] or 0) / 100])
    g = payload["grand"]
    ws.append(["Total", g["lines"], _money(g["equipment_cost_ext"]), _money(g["equipment_sell_ext"]), float(g["labor_hours"]), _money(g["labor_cost"]),
               _money(g["labor_sell"]), _money(g["total_cost"]), _money(g["total_sell"]), _money(g["profit"]), (g["margin"] or 0) / 100])
    for c in ws[ws.max_row]:
        c.font = bold
    for row in ws.iter_rows(min_row=8, min_col=3, max_col=10):
        for c in row:
            c.number_format = '"$"#,##0.00'
    for row in ws.iter_rows(min_row=8, min_col=11, max_col=11):
        for c in row:
            c.number_format = "0.0%"
    ws.column_dimensions["A"].width = 34
    for i in range(2, 12):
        ws.column_dimensions[get_column_letter(i)].width = 15
    if payload["warnings"]:
        ws.append([])
        ws.append(["Estimator Notes"])
        for w in payload["warnings"]:
            ws.append([w["text"]])
    used = {"Summary"}
    labor_hdr = []
    for t in rules.LABOR_TYPES:
        labor_hdr += ["%s Units" % t["label"], "%s Cost" % t["label"], "%s Sell" % t["label"]]
    for r in payload["rooms"]:
        s = wb.create_sheet(rules.sheet_name(r["name"], used))
        s.append([r["name"]])
        s["A1"].font = Font(bold=True, size=13)
        s.append(["Area", "Item", "Manufacturer", "Model #", "Description", "Qty", "Dealer Cost Ea", "Equipment Cost Ext", "Markup", "Sell Ea",
                  "Equipment Sell Ext"] + labor_hdr + ["Total Cost", "Total Sell", "Profit", "Margin", "Source"])
        for c in s[2]:
            c.font, c.fill, c.alignment = bold, head_fill, Alignment(wrap_text=True, vertical="top")
        for ln in r["lines"]:
            t = ln["totals"]
            row = [ln["area"], "Note" if ln["kind"] == "note" else rules.item_label(ln["category"], ln["manufacturer"], ln["part"], ln["description"]),
                   ln["manufacturer"], ln["part"], ln["description"], ln["qty"], _money(ln["cost"]), _money(t["equipment_cost_ext"]),
                   float(ln["markup"]) if ln["markup"] is not None else None, _money(ln["sell"]), _money(t["equipment_sell_ext"])]
            for p in t["labor"]:
                row += [float(p["hours"]), _money(p["cost"]), _money(p["sell"])]
            row += [_money(t["total_cost"]), _money(t["total_sell"]), _money(t["profit"]), (t["margin"] or 0) / 100 if t["total_sell"] else None,
                    ln.get("source_name") or ""]
            s.append(row)
        tt = r["totals"]
        tot = ["Room total", "", "", "", "", tt["lines"], "", _money(tt["equipment_cost_ext"]), "", "", _money(tt["equipment_sell_ext"])]
        for t in rules.LABOR_TYPES:
            tot += [float(tt["hours_by_type"].get(t["id"], 0)), "", ""]
        tot += [_money(tt["total_cost"]), _money(tt["total_sell"]), _money(tt["profit"]), (tt["margin"] or 0) / 100 if tt["total_sell"] else None, ""]
        s.append(tot)
        for c in s[s.max_row]:
            c.font = bold
        s.freeze_panes = "F3"
        s.column_dimensions["A"].width, s.column_dimensions["C"].width, s.column_dimensions["D"].width, s.column_dimensions["E"].width = 18, 18, 20, 44
        ncol = 11 + 3 * len(rules.LABOR_TYPES) + 5
        for i in range(6, ncol + 1):
            s.column_dimensions[get_column_letter(i)].width = 12
        for row in s.iter_rows(min_row=3, min_col=7, max_col=ncol - 1):
            for c in row:
                if isinstance(c.value, float):
                    c.number_format = '"$"#,##0.00'
        for row in s.iter_rows(min_row=3, min_col=ncol - 1, max_col=ncol - 1):
            for c in row:
                c.number_format = "0.0%"
        for row in s.iter_rows(min_row=3, min_col=9, max_col=9):
            for c in row:
                c.number_format = "0.000"
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def export_csv(est, payload, room_id=None):
    """Detail CSV: every room (Room column first) or one room — the dashboard's 71-column layout with a Room column and
    the source date, one row per line, room totals rows."""
    buf = io.StringIO()
    w = csv.writer(buf)
    labor_hdr = []
    for t in rules.LABOR_TYPES:
        labor_hdr += ["%s Cost" % t["label"], "%s Cost Extended" % t["label"], "%s Owner up" % t["label"], "%s Sell" % t["label"], "%s Sell Extended" % t["label"]]
    w.writerow(["Room", "Area", "Item", "Manufacturer", "Model #", "Description", "Qty", "Dealer Cost Ea", "Equipment Cost Ext", "Markup", "Sell Ea",
                "Equipment Sell Ext"] + labor_hdr + ["Total Cost", "Total Sell", "Profit", "Margin", "Source Date", "Source"])
    for r in payload["rooms"]:
        if room_id and r["id"] != room_id:
            continue
        for ln in r["lines"]:
            t = ln["totals"]
            row = [r["name"], ln["area"], "Note" if ln["kind"] == "note" else rules.item_label(ln["category"], ln["manufacturer"], ln["part"], ln["description"]),
                   ln["manufacturer"], ln["part"], ln["description"], ln["qty"], "%.2f" % ln["cost"], "%.2f" % t["equipment_cost_ext"],
                   "%.4f" % ln["markup"] if ln["markup"] is not None else "", "%.2f" % ln["sell"], "%.2f" % t["equipment_sell_ext"]]
            for p in t["labor"]:
                rc, rs = rules.dec(p["rate_cost"]), rules.dec(p["rate_sell"])
                row += ["%.2f" % rc, "%.2f" % p["cost"], "%.3f" % (rs / rc) if rc else "", "%.2f" % rs, "%.2f" % p["sell"]]
            row += ["%.2f" % t["total_cost"], "%.2f" % t["total_sell"], "%.2f" % t["profit"], "%.1f%%" % t["margin"] if t["margin"] is not None else "",
                    ln.get("date_label") or rules.date_key_label(ln.get("date_key")), ln.get("source_name") or ""]
            w.writerow(row)
        tt = r["totals"]
        w.writerow([r["name"], "ROOM TOTAL", "", "", "", "", tt["lines"], "", "%.2f" % tt["equipment_cost_ext"], "", "", "%.2f" % tt["equipment_sell_ext"]]
                   + [""] * (5 * len(rules.LABOR_TYPES)) + ["%.2f" % tt["total_cost"], "%.2f" % tt["total_sell"], "%.2f" % tt["profit"],
                                                            "%.1f%%" % tt["margin"] if tt["margin"] is not None else "", "", ""])
    return buf.getvalue()


def export_text(est, payload):
    """Clipboard text: one block per line + totals (the dashboard's Copy to Clipboard)."""
    out = ["%s%s" % (est.title or "PACE Estimate", (" — " + est.client_name) if est.client_name else "")]
    for r in payload["rooms"]:
        out.append("")
        out.append("== %s (%d items · sell $%s)" % (r["name"], r["totals"]["lines"], "{:,.2f}".format(rules.dec(r["totals"]["total_sell"]))))
        for ln in r["lines"]:
            t = ln["totals"]
            out.append("%s %s — %s | Qty %d | Product Cost $%s | Sell Ea $%s | Labor Sell $%s | Profit $%s" % (
                ln["manufacturer"], ln["part"], ln["description"], ln["qty"], "{:,.2f}".format(rules.dec(ln["cost"])), "{:,.2f}".format(rules.dec(ln["sell"])),
                "{:,.2f}".format(rules.dec(t["labor_sell"])), "{:,.2f}".format(rules.dec(t["profit"]))))
    g = payload["grand"]
    out.append("")
    out.append("TOTAL: Cost $%s | Sell $%s | Profit $%s | Margin %s | Labor %s h" % (
        "{:,.2f}".format(rules.dec(g["total_cost"])), "{:,.2f}".format(rules.dec(g["total_sell"])), "{:,.2f}".format(rules.dec(g["profit"])),
        ("%.1f%%" % g["margin"]) if g["margin"] is not None else "—", "{:,.1f}".format(rules.dec(g["labor_hours"]))))
    return "\n".join(out)
