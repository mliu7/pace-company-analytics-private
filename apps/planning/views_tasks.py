"""Production › Active tasks & PM report (SharePoint spec §6.2; parity A.7 AT-01…AT-09) — a view over StatusRow."""

import io
from datetime import datetime

from django.http import HttpResponse
from django.shortcuts import render
from django.utils import timezone

from . import rules
from .models import StatusRow
from .views import _pctx, _today, status_dataset, _user_name

PM_REPORT_TITLE = "PM Active Task Report"
PACE_ADDRESS = ["Pace Systems, Inc.", "2040 Corporate Lane", "Naperville, IL 60563"]


def _active_qs():
    return StatusRow.objects.filter(completed=False, percent__lt=100)


def _tasks_for(pm=None, division=None):
    """Active tasks as plain dicts with the priority bucket; `pm` matches PM1 or PM2 (name or raw, case-insensitive)."""
    qs = _active_qs().select_related("pm", "pm2", "project")
    if division:
        qs = qs.filter(division=division)
    today = _today()
    out = []
    key = (pm or "").strip().upper()
    for r in qs.order_by("division", "end", "name"):
        names = {(r.pm.canonical_name if r.pm else r.pm_raw).upper(), (r.pm2.canonical_name if r.pm2 else r.pm2_raw).upper()}
        if key and key not in names and key != "ALL":
            continue
        out.append({
            "id": r.id, "division": r.division, "name": r.name, "project_number": r.project_number_raw or r.job_key,
            "pm": r.pm.canonical_name if r.pm else r.pm_raw, "pm2": r.pm2.canonical_name if r.pm2 else r.pm2_raw,
            "start": r.start, "end": r.end, "hours_left": float(r.hours_left or 0), "percent": float(r.percent or 0),
            "phase_status": r.phase_status, "equipment_status": r.equipment_status,
            "priority": rules.priority_bucket(r.end, today), "days_to_stop": rules.days_until(r.end, today),
            "priority_label": rules.PRIORITY_LABELS[rules.priority_bucket(r.end, today)],
        })
    return out


def tasks(request):
    ds = status_dataset(_active_qs())
    can_write = bool(getattr(request, "acc", None) and request.acc.planning_write)
    return render(request, "planning/tasks.html", _pctx(request, "tasks", dataset=ds, can_write=can_write))


def tasks_report(request):
    """Print view of the PM report (AT-07): letterhead, six cards, schedule totals, the table across 040 / 070 / 080."""
    pm = request.GET.get("pm") or ""
    rows = _tasks_for(pm)
    today = _today()
    summary = rules.task_summary(rows, today)
    per_div = {d: rules.task_summary([r for r in rows if r["division"] == d], today) for d in sorted({r["division"] for r in rows} | {"040", "070", "080"})}
    order = {"red": 0, "yellow": 1, "green": 2, "none": 3}
    attention = sorted([r for r in rows if r["priority"] in ("red", "yellow", "none")], key=lambda r: (order[r["priority"]], r["end"] or datetime.max.date(), r["name"]))
    return render(request, "planning/tasks_report.html", _pctx(request, "tasks", pm=pm or "All project managers", rows=rows, summary=summary,
                                                               per_div=per_div, attention=attention, generated=timezone.localtime(),
                                                               address=PACE_ADDRESS, title=PM_REPORT_TITLE))


def tasks_export(request):
    """4-sheet styled Excel (AT-07): PM Report (letterhead + workload + dates / status + attention + all active with
    live SUM / AVERAGE / COUNTIF), Summary, Attention Required, All Active Tasks — openpyxl, logo on sheets 1 and 2."""
    import openpyxl
    from openpyxl.drawing.image import Image as XLImage
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter

    pm = request.GET.get("pm") or ""
    rows = _tasks_for(pm)
    today = _today()
    summary = rules.task_summary(rows, today)
    label = pm or "All PMs"
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "PM Report"
    blue = PatternFill("solid", fgColor="1F4E79")
    yellow = PatternFill("solid", fgColor="FFF2CC")
    alt = PatternFill("solid", fgColor="F2F2F2")
    fill_red = PatternFill("solid", fgColor="F8CBAD")
    fill_org = PatternFill("solid", fgColor="FFE699")
    fill_grn = PatternFill("solid", fgColor="C6E0B4")
    fill_gry = PatternFill("solid", fgColor="D9D9D9")
    white_b = Font(bold=True, color="FFFFFF")
    bold = Font(bold=True)
    thin = Side(style="thin", color="BFBFBF")
    box = Border(left=thin, right=thin, top=thin, bottom=thin)
    logo_path = None
    try:
        from django.contrib.staticfiles import finders
        logo_path = finders.find("dashboard/pace-logo-192.png")
    except Exception:  # noqa
        logo_path = None

    def logo(sheet, anchor):
        if logo_path:
            try:
                img = XLImage(logo_path)
                img.width, img.height = 72, 72
                sheet.add_image(img, anchor)
            except Exception:  # noqa
                pass

    def section(sheet, row, text, ncols, fill=blue, font=white_b):
        sheet.cell(row=row, column=1, value=text).font = font
        for c in range(1, ncols + 1):
            sheet.cell(row=row, column=c).fill = fill

    # ---- letterhead (rows 1-11)
    logo(ws, "A1")
    ws["C1"] = PM_REPORT_TITLE; ws["C1"].font = Font(bold=True, size=16)
    ws["C2"] = "Project Manager:"; ws["D2"] = label; ws["C2"].font = bold
    ws["C3"] = "Report Scope:"; ws["D3"] = "Active tasks across schedules 040 / 070 / 080"; ws["C3"].font = bold
    ws["C4"] = "Generated:"; ws["D4"] = timezone.localtime().strftime("%Y-%m-%d %H:%M"); ws["C4"].font = bold
    for i, line in enumerate(PACE_ADDRESS):
        ws.cell(row=6 + i, column=3, value=line)
    ws["C10"] = "Description of Request:"; ws["C10"].font = bold
    ws["D10"] = "Every open Master Schedule task for the project manager with union hours remaining, critical stop and both status pills."
    cols = ["Item", "Schedule", "Project #", "Project Manager (PM1 / PM2)", "Description", "Union Hours", "Complete", "Days to Stop",
            "Start", "Critical Stop", "Priority", "Status (S1 / S2)"]
    section(ws, 12, "PM Active Task Report and Total Counts", len(cols))
    ws.cell(row=12, column=6, value="Workload").font = white_b
    ws.cell(row=12, column=9, value="Dates / Status").font = white_b
    for c, h in enumerate(cols, 1):
        cell = ws.cell(row=13, column=c, value=h)
        cell.font = bold; cell.fill = yellow; cell.border = box; cell.alignment = Alignment(wrap_text=True, vertical="center")
    r0 = 14
    for i, t in enumerate(rows):
        r = r0 + i
        pmtxt = t["pm"] + (" / " + t["pm2"] if t["pm2"] else "")
        st = " / ".join(x for x in (t["phase_status"], t["equipment_status"]) if x)
        vals = [i + 1, t["division"], t["project_number"], pmtxt, t["name"], t["hours_left"], t["percent"] / 100.0,
                t["days_to_stop"], t["start"], t["end"], t["priority_label"], st]
        for c, v in enumerate(vals, 1):
            cell = ws.cell(row=r, column=c, value=v)
            cell.border = box
            if i % 2:
                cell.fill = alt
        ws.cell(row=r, column=7).number_format = "0%"
        for c in (9, 10):
            ws.cell(row=r, column=c).number_format = "yyyy-mm-dd"
        f = {"red": fill_red, "yellow": fill_org, "green": fill_grn, "none": fill_gry}[t["priority"]]
        ws.cell(row=r, column=11).fill = f
        ws.cell(row=r, column=8).fill = f
        if t["priority"] == "red":
            ws.cell(row=r, column=10).fill = f
    rN = r0 + max(len(rows), 1) - 1
    tot = rN + 1
    ws.cell(row=tot, column=5, value="Subtotal (live formulas)").font = bold
    ws.cell(row=tot, column=6, value="=SUM(F%d:F%d)" % (r0, rN)).font = bold
    ws.cell(row=tot, column=7, value="=IFERROR(AVERAGE(G%d:G%d),0)" % (r0, rN)).font = bold
    ws.cell(row=tot, column=7).number_format = "0%"
    ws.cell(row=tot, column=11, value='=COUNTIF(K%d:K%d,"%s")' % (r0, rN, rules.PRIORITY_LABELS["red"])).font = bold
    ws.cell(row=tot, column=12, value="past due").font = bold
    # ---- summary blocks
    r = tot + 2
    section(ws, r, "Original PM Report Summary", len(cols)); r += 1
    for k, lab in (("active", "Active tasks"), ("union_hours", "Union hours"), ("past_due", "Past due"), ("due_soon", "Due within 14 days"),
                   ("missing", "Missing stop date"), ("avg_pct", "Average complete (%)")):
        ws.cell(row=r, column=1, value=lab).font = bold; ws.cell(row=r, column=3, value=summary[k]); r += 1
    r += 1
    section(ws, r, "Schedule Breakdown", len(cols)); r += 1
    for h_i, h in enumerate(("Schedule", "Active tasks", "Union hours", "Past due", "Due soon", "Average complete"), 1):
        ws.cell(row=r, column=h_i, value=h).font = bold; ws.cell(row=r, column=h_i).fill = yellow
    r += 1
    for d in ("040", "070", "080"):
        s = rules.task_summary([t for t in rows if t["division"] == d], today)
        for c_i, v in enumerate((d, s["active"], s["union_hours"], s["past_due"], s["due_soon"], s["avg_pct"]), 1):
            ws.cell(row=r, column=c_i, value=v)
        r += 1
    r += 1
    section(ws, r, "Attention Required", len(cols)); r += 1
    ws.cell(row=r, column=1, value="Past critical stop: %d · Due within 14 days: %d · Missing stop date: %d · Urgent union hours: %.0f" % (
        summary["past_due"], summary["due_soon"], summary["missing"], summary["urgent_hours"])); r += 2
    ws.cell(row=r, column=1, value="TOTAL ACTIVE TASKS").font = bold; ws.cell(row=r, column=3, value=summary["active"]).font = bold; r += 1
    ws.cell(row=r, column=1, value="TOTAL UNION HOURS").font = bold; ws.cell(row=r, column=3, value=summary["union_hours"]).font = bold
    widths = [6, 9, 11, 26, 52, 12, 10, 12, 12, 13, 18, 40]
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = "A14"
    ws.auto_filter.ref = "A13:%s%d" % (get_column_letter(len(cols)), rN)
    ws.page_setup.orientation = "landscape"; ws.page_setup.fitToWidth = 1; ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.print_title_rows = "13:13"

    # ---- Summary sheet
    s2 = wb.create_sheet("Summary")
    logo(s2, "A1")
    s2["C1"] = "%s — %s" % (PM_REPORT_TITLE, label); s2["C1"].font = Font(bold=True, size=14)
    kp = [("Active tasks", summary["active"]), ("Union hours", summary["union_hours"]), ("Past due", summary["past_due"]),
          ("Due within 14 days", summary["due_soon"]), ("Missing stop date", summary["missing"])]
    for i, (lab, v) in enumerate(kp):
        c = s2.cell(row=4, column=1 + i * 2, value=lab); c.font = bold; c.fill = yellow
        c = s2.cell(row=5, column=1 + i * 2, value=v); c.font = Font(bold=True, size=18)
    section(s2, 7, "Schedule Workload", 6); r = 8
    for h_i, h in enumerate(("Schedule", "Active Tasks", "Union Hours", "Past Due", "Due Soon", "Average Complete"), 1):
        s2.cell(row=r, column=h_i, value=h).font = bold
    r += 1
    for d in ("040", "070", "080"):
        s = rules.task_summary([t for t in rows if t["division"] == d], today)
        for c_i, v in enumerate((d, s["active"], s["union_hours"], s["past_due"], s["due_soon"], s["avg_pct"]), 1):
            s2.cell(row=r, column=c_i, value=v)
        r += 1
    r += 1
    section(s2, r, "Report Sections", 6); r += 1
    for name in ("PM Report", "Summary", "Attention Required", "All Active Tasks"):
        s2.cell(row=r, column=1, value=name); r += 1
    for i, w in enumerate((22, 12, 22, 12, 22, 12), 1):
        s2.column_dimensions[get_column_letter(i)].width = w

    # ---- detail sheets
    detail_cols = ["Item", "Schedule", "Project Number", "Project / Task", "PM1", "PM2", "Start", "Critical Stop", "Union Hours Remaining",
                   "Percent Complete", "Days to Stop", "Priority", "Status", "Status V2"]

    def detail(sheet, data):
        for c, h in enumerate(detail_cols, 1):
            cell = sheet.cell(row=1, column=c, value=h); cell.font = bold; cell.fill = yellow; cell.border = box
        for i, t in enumerate(data):
            vals = [i + 1, t["division"], t["project_number"], t["name"], t["pm"], t["pm2"], t["start"], t["end"], t["hours_left"],
                    t["percent"] / 100.0, t["days_to_stop"], t["priority_label"], t["phase_status"], t["equipment_status"]]
            for c, v in enumerate(vals, 1):
                cell = sheet.cell(row=2 + i, column=c, value=v); cell.border = box
            sheet.cell(row=2 + i, column=10).number_format = "0%"
            for c in (7, 8):
                sheet.cell(row=2 + i, column=c).number_format = "yyyy-mm-dd"
            sheet.cell(row=2 + i, column=12).fill = {"red": fill_red, "yellow": fill_org, "green": fill_grn, "none": fill_gry}[t["priority"]]
        n = len(data)
        sheet.cell(row=n + 2, column=4, value="Section totals").font = bold
        sheet.cell(row=n + 2, column=9, value="=SUM(I2:I%d)" % (n + 1)).font = bold
        sheet.cell(row=n + 2, column=10, value="=IFERROR(AVERAGE(J2:J%d),0)" % (n + 1)).font = bold
        sheet.cell(row=n + 2, column=10).number_format = "0%"
        for i, w in enumerate((6, 9, 14, 50, 18, 18, 12, 13, 12, 10, 11, 20, 30, 30), 1):
            sheet.column_dimensions[get_column_letter(i)].width = w
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = "A1:%s%d" % (get_column_letter(len(detail_cols)), max(n + 1, 2))

    order = {"red": 0, "yellow": 1, "none": 2, "green": 3}
    att = sorted([t for t in rows if t["priority"] != "green"], key=lambda t: (order[t["priority"]], t["end"] or datetime.max.date()))
    detail(wb.create_sheet("Attention Required"), att)
    detail(wb.create_sheet("All Active Tasks"), rows)

    buf = io.BytesIO()
    wb.save(buf)
    fname = "PACE_PM_Report_%s_%s.xlsx" % ((pm or "All").replace(" ", "_").replace("/", "-"), today.isoformat())
    resp = HttpResponse(buf.getvalue(), content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
    resp["Content-Disposition"] = 'attachment; filename="%s"' % fname
    return resp
