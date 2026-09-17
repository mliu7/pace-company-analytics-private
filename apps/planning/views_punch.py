"""Production › Punch lists (SharePoint spec §6.4; parity A.8 PU-01…PU-08)."""

import csv
import io
import re

from django.conf import settings
from django.db.models import Prefetch
from django.http import HttpResponse, JsonResponse
from django.shortcuts import get_object_or_404, render
from django.utils import timezone
from django.views.decorators.http import require_POST

from apps.access.audit import log as audit
from apps.core.models import Employee

from . import loaders, rules
from .models import PunchDepartment, PunchHistory, PunchItem, PunchProject, Source, StatusRow
from .views import DIVISIONS, _bad, _body, _iso, _pctx, _pm_roster, _resolve_employee, _resolve_project, _today, _user, _user_name

ITEM_FIELDS = ("active", "date_entered", "description", "bic", "critical", "due_by", "date_completed", "assigned", "engineer_signoff", "verified")


def _division(request):
    d = request.GET.get("div") or request.POST.get("div") or ""
    if d not in DIVISIONS:
        first = PunchProject.objects.order_by("division").values_list("division", flat=True).first()
        d = first if first in DIVISIONS else (settings.MODELLED_DIVISION_CODE if settings.MODELLED_DIVISION_CODE in DIVISIONS else "040")
    return d


def _item_dict(it):
    return {"id": it.id, "project_id": it.punch_project_id, "position": it.position, "active": it.active,
            "date_entered": _iso(it.date_entered), "description": it.description, "bic": it.bic, "critical": it.critical,
            "due_by": _iso(it.due_by), "date_completed": _iso(it.date_completed), "assigned": it.assigned,
            "engineer_signoff": it.engineer_signoff, "verified": it.verified, "version": it.version,
            "updated_at": _iso(it.updated_at), "updated_by": _user_name(it.updated_by),
            "attachments": [{"id": a.id, "name": a.name, "size": a.size} for a in getattr(it, "_atts", [])]}


def _project_dict(p, status_by_key):
    src = status_by_key.get(rules.code_digits(p.code)[:6])
    return {"id": p.id, "code": p.code, "title": p.title or p.code, "pm_id": p.pm_id, "pm": p.pm.canonical_name if p.pm else (p.pm_raw or "Unassigned"),
            "project_id": p.project_id, "project_number": p.project.canonical_project_number if p.project else "",
            "project_url": "/projects/%s/" % p.project.canonical_project_number if p.project else "",
            "source": p.source, "items": [_item_dict(i) for i in p.items.all()],
            "src": src}


def punch_dataset(division):
    projects = list(PunchProject.objects.filter(division=division).select_related("pm", "project")
                    .prefetch_related(Prefetch("items", queryset=PunchItem.objects.select_related("updated_by").order_by("position", "id"))))
    # planner chips: status / status2 / stop / hours from the status board rows of the same job (Σ hours over its rows)
    status_by_key = {}
    for s in StatusRow.objects.filter(division=division, completed=False).order_by("id"):
        k = (s.project_number_raw or s.job_key)[:6]
        if not k:
            continue
        cur = status_by_key.setdefault(k, {"status": s.phase_status, "status2": s.equipment_status, "finish": _iso(s.end), "start": _iso(s.start),
                                           "hours": 0.0, "pct": float(s.percent or 0), "row_id": s.id})
        cur["hours"] += float(s.hours_left or 0)
        if s.end and (not cur["finish"] or _iso(s.end) > cur["finish"]):
            cur["finish"] = _iso(s.end)
    roster = [p for p in _pm_roster() if p["div"] == division or p["div"] == ""]
    referenced = {p.pm_id for p in projects if p.pm_id}
    roster_ids = {p["id"] for p in roster}
    for p in _pm_roster():
        if p["id"] in referenced and p["id"] not in roster_ids:
            roster.append(p)
    return {"division": division, "divisions": DIVISIONS, "today": _today().isoformat(),
            "projects": [_project_dict(p, status_by_key) for p in projects],
            "pms": sorted(roster, key=lambda p: p["name"]), "bics": PunchDepartment.names(),
            "critical": {str(k): {"label": v, "color": rules.CRITICAL_COLORS[k]} for k, v in rules.CRITICAL_LABELS.items()}}


def punch(request, pk=None):
    division = _division(request)
    open_pk = pk
    if pk is not None:
        pp = get_object_or_404(PunchProject, pk=pk)
        division = pp.division
    ds = punch_dataset(division)
    can_write = bool(getattr(request, "acc", None) and request.acc.planning_write)
    if request.GET.get("format") == "json":
        return JsonResponse(ds)
    return render(request, "planning/punch.html", _pctx(request, "punch", dataset=ds, division=division, open_pk=open_pk, can_write=can_write))


def punch_project(request, pk):
    return punch(request, pk=pk)


def _hist(pp, item, user, field, old, new):
    PunchHistory.objects.create(punch_project=pp, item=item, changed_by=user, field=field, old=loaders._s(old), new=loaders._s(new))


def _set_item_field(it, field, value):
    if field == "active":
        return bool(value) if not isinstance(value, str) else value.lower() in ("1", "true", "on", "yes")
    if field in ("date_entered", "due_by", "date_completed"):
        return rules.parse_date_cell(value)
    if field == "critical":
        try:
            v = int(value) if value not in (None, "") else None
        except (TypeError, ValueError):
            v = None
        return v if v in (1, 2, 3, 4, 5) else None
    if field == "description":
        return value or ""
    return (value or "").strip()[:200 if field != "bic" else 60]


@require_POST
def punch_edit(request):
    """{op: item_update|item_add|item_delete|project_add|project_update|project_delete|seed, ...}"""
    b = _body(request)
    op = b.get("op") or ""
    user = _user(request)
    if op == "seed":
        division = b.get("div") if b.get("div") in DIVISIONS else _division(request)
        n = loaders.seed_punch_projects_from_status(division, user=user)
        audit("write_action", request, meta_op="punch_seed", meta_division=division, meta_added=n)
        return JsonResponse({"ok": True, "added": n, "dataset": punch_dataset(division)})
    if op == "project_add":
        division = b.get("div") if b.get("div") in DIVISIONS else _division(request)
        code = rules.punch_code((b.get("code") or "").strip())
        if not code:
            return _bad("Project code is required.")
        if PunchProject.objects.filter(division=division, code__iexact=code).exists():
            return _bad("A project with code %s already exists in %s." % (code, division))
        pid, _c = _resolve_project(rules.code_digits(code))
        pm_name = (b.get("pm") or "").strip()
        pm_id = _resolve_employee(pm_name) if pm_name and pm_name.lower() != "unassigned" else None
        pp = PunchProject.objects.create(division=division, code=code, title=(b.get("title") or code).strip()[:300], project_id=pid,
                                         pm_id=pm_id, pm_raw=(Employee.objects.get(pk=pm_id).canonical_name if pm_id else "Unassigned")[:120],
                                         source=Source.MANUAL, created_by=user, updated_by=user)
        _hist(pp, None, user, "created", "", pp.code)
        audit("write_action", request, meta_op="punch_project_add", meta_project=pp.id)
        return JsonResponse({"ok": True, "project": _project_dict(pp, {})})
    if op == "item_add":
        pp = PunchProject.objects.filter(pk=b.get("project_id")).first()
        if pp is None:
            return _bad("project not found", 404)
        pos = (pp.items.order_by("-position").values_list("position", flat=True).first() or 0) + 1
        it = PunchItem.objects.create(punch_project=pp, position=pos, active=False, bic=rules.PUNCH_ITEM_DEFAULTS["bic"],
                                      critical=rules.PUNCH_ITEM_DEFAULTS["critical"], date_entered=_today(), created_by=user, updated_by=user)
        _hist(pp, it, user, "item_added", "", str(it.id))
        pp.updated_by = user; pp.save(update_fields=["updated_by", "updated_at"])
        audit("write_action", request, meta_op="punch_item_add", meta_item=it.id)
        return JsonResponse({"ok": True, "item": _item_dict(it)})
    if op == "item_update":
        it = PunchItem.objects.select_related("punch_project", "updated_by").filter(pk=b.get("id")).first()
        if it is None:
            return _bad("item not found", 404)
        try:
            version = int(b.get("version") or 0)
        except (TypeError, ValueError):
            version = 0
        if version and version != it.version:
            return _bad("changed by %s at %s — reload" % (_user_name(it.updated_by) or "someone", timezone.localtime(it.updated_at).strftime("%H:%M")),
                        status=409, item=_item_dict(it))
        fields = b.get("fields") or ({b.get("field"): b.get("value")} if b.get("field") else {})
        changed = []
        for f, v in fields.items():
            if f not in ITEM_FIELDS:
                return _bad("field %r is not editable" % f)
            new = _set_item_field(it, f, v)
            old = getattr(it, f)
            if old != new:
                setattr(it, f, new)
                changed.append((f, old, new))
        if changed:
            it.updated_by = user
            it.version += 1
            it.save()
            for f, o, n in changed:
                _hist(it.punch_project, it, user, f, o, n)
            audit("write_action", request, meta_op="punch_item_update", meta_item=it.id, meta_fields=[c[0] for c in changed])
        return JsonResponse({"ok": True, "item": _item_dict(it), "changed": [c[0] for c in changed]})
    if op == "item_delete":
        it = PunchItem.objects.select_related("punch_project").filter(pk=b.get("id")).first()
        if it is None:
            return _bad("item not found", 404)
        _hist(it.punch_project, None, user, "item_deleted", (it.description or "")[:200], "")
        audit("write_action", request, meta_op="punch_item_delete", meta_item=it.id)
        it.delete()
        return JsonResponse({"ok": True, "deleted": b.get("id")})
    if op in ("project_update", "project_delete"):
        pp = PunchProject.objects.select_related("pm", "project").filter(pk=b.get("id")).first()
        if pp is None:
            return _bad("project not found", 404)
        if op == "project_delete":
            n = pp.items.count()
            audit("write_action", request, meta_op="punch_project_delete", meta_project=pp.id, meta_code=pp.code, meta_items=n)
            pp.delete()
            return JsonResponse({"ok": True, "deleted": b.get("id")})
        fields = b.get("fields") or {}
        if "title" in fields:
            t = (fields["title"] or "").strip()[:300]
            if t and t != pp.title:
                _hist(pp, None, user, "title", pp.title, t); pp.title = t
        if "pm" in fields:
            name = (fields["pm"] or "").strip()
            pm_id = _resolve_employee(name) if name and name.lower() != "unassigned" else None
            raw = Employee.objects.get(pk=pm_id).canonical_name if pm_id else "Unassigned"
            if pm_id != pp.pm_id or raw != pp.pm_raw:
                _hist(pp, None, user, "pm", pp.pm_raw, raw); pp.pm_id, pp.pm_raw = pm_id, raw[:120]
        if "code" in fields:
            c = rules.punch_code((fields["code"] or "").strip())
            if c and c != pp.code:
                _hist(pp, None, user, "code", pp.code, c); pp.code = c
                pp.project_id, _c = _resolve_project(rules.code_digits(c))
        pp.updated_by = user
        pp.save()
        pp = PunchProject.objects.select_related("pm", "project").get(pk=pp.pk)
        audit("write_action", request, meta_op="punch_project_update", meta_project=pp.id)
        return JsonResponse({"ok": True, "project": _project_dict(pp, {})})
    return _bad("unknown op %r" % op)


HEADERS = ["ACTIVE", "DATE ENTERED", "DESCRIPTION OF ITEM", "BIC", "CRITICAL (1-5)", "DUE BY", "DATE COMPLETED", "ASSIGNED",
           "ENGINEER SIGNOFF", "VERIFIED (INITIALS/DATE)"]


def punch_export(request):
    """Excel: one sheet per project (title row + header row + items); CSV: one file with PROJECT CODE / TITLE / PM columns."""
    division = _division(request)
    fmt = request.GET.get("fmt") or "xlsx"
    projects = list(PunchProject.objects.filter(division=division).select_related("pm").prefetch_related("items").order_by("code"))
    ids = request.GET.get("ids")
    if ids:
        want = {int(x) for x in ids.split(",") if x.isdigit()}
        projects = [p for p in projects if p.id in want]
    today = _today()
    if fmt == "csv":
        buf = io.StringIO()
        w = csv.writer(buf, lineterminator="\r\n")
        w.writerow(["PROJECT CODE", "PROJECT TITLE", "PM"] + HEADERS)
        for p in projects:
            pm = p.pm.canonical_name if p.pm else p.pm_raw
            items = list(p.items.all())
            if not items:
                w.writerow([p.code, p.title, pm] + [""] * len(HEADERS))
            for it in items:
                w.writerow([p.code, p.title, pm, "TRUE" if it.active else "FALSE", _iso(it.date_entered) or "", it.description, it.bic,
                            it.critical or "", _iso(it.due_by) or "", _iso(it.date_completed) or "", it.assigned, it.engineer_signoff, it.verified])
        resp = HttpResponse("﻿" + buf.getvalue(), content_type="text/csv; charset=utf-8")
        resp["Content-Disposition"] = 'attachment; filename="%s_Punch_Lists_Export_%s.csv"' % (division, today.isoformat())
        return resp
    import openpyxl
    from openpyxl.styles import Font, PatternFill
    from openpyxl.utils import get_column_letter
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    used = set()
    for p in projects:
        name = re.sub(r"[\\/?*\[\]:]", "-", p.code)[:31] or "Project"
        base, n = name, 1
        while name in used:
            n += 1
            name = ("%s_%d" % (base[:28], n))[:31]
        used.add(name)
        ws = wb.create_sheet(name)
        ws.cell(row=1, column=1, value=p.title or p.code).font = Font(bold=True, size=13)
        for c, h in enumerate(HEADERS, 1):
            cell = ws.cell(row=2, column=c, value=h); cell.font = Font(bold=True); cell.fill = PatternFill("solid", fgColor="FFF2CC")
        for i, it in enumerate(p.items.all()):
            vals = ["TRUE" if it.active else "FALSE", it.date_entered, it.description, it.bic, it.critical, it.due_by, it.date_completed,
                    it.assigned, it.engineer_signoff, it.verified]
            for c, v in enumerate(vals, 1):
                cell = ws.cell(row=3 + i, column=c, value=v)
                if c in (2, 6, 7):
                    cell.number_format = "yyyy-mm-dd"
        for i, w in enumerate((8, 13, 60, 18, 12, 12, 14, 16, 18, 22), 1):
            ws.column_dimensions[get_column_letter(i)].width = w
        ws.freeze_panes = "A3"
    if not wb.sheetnames:
        wb.create_sheet("Empty").cell(row=1, column=1, value="No projects in %s" % division)
    buf = io.BytesIO()
    wb.save(buf)
    resp = HttpResponse(buf.getvalue(), content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
    resp["Content-Disposition"] = 'attachment; filename="%s_Punch_Lists_Export_%s.xlsx"' % (division, today.isoformat())
    return resp
