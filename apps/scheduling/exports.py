"""Resource Scheduler exports (RS-13): the week CSV exactly as the dashboard wrote it, the per-resource weekly print view
and an XLSX workbook. JSON import / reset / P-drive save are gone by design — the database is the record."""

import csv
import io
from collections import defaultdict
from datetime import timedelta

from django.http import HttpResponse
from django.template.loader import render_to_string

from . import maths as M
from .services import iso


def _day_label(d):
    return "%s %s %d" % (d.strftime("%a"), d.strftime("%b"), d.day)


def _cell_text(items):
    """'Project • Trim 8h, Rough 4h' per project, one project per line."""
    by_proj = defaultdict(list)
    for a in items:
        by_proj[a["project_name"]].append(a)
    lines = []
    for pn, arr in sorted(by_proj.items()):
        parts = ["%s %sh" % (a["phase_name"], M.h(a["hours"])) if a["phase_name"] else "%sh" % M.h(a["hours"]) for a in arr]
        lines.append("%s • %s" % (pn, ", ".join(parts)))
    return "\n".join(lines)


def week_tables(state, monday):
    """The rows behind both the CSV and the XLSX."""
    days = M.week_days(monday)
    A = [a for a in state["assignments"] if a["date"] and days[0] <= a["date"] <= days[6]]
    rs_by, ps_by = state["resources_by_id"], state["projects_by_id"]
    projects_row = []
    for d in days:
        names = sorted({a["project_name"] for a in A if a["date"] == d})
        projects_row.append("\n".join(names))
    rows = []
    for r in sorted(state["resources"], key=lambda x: x["sort_name"].lower()):
        cells = []
        for d in days:
            if M.on_pto(r, d):
                cells.append("PTO / Vacation")
            else:
                cells.append(_cell_text([a for a in A if a["resource_id"] == r["id"] and a["date"] == d]))
        rows.append((r, cells))
    by_res = [(r["name"], r["trade"], sum(a["hours"] for a in A if a["resource_id"] == r["id"])) for r in sorted(state["resources"], key=lambda x: x["sort_name"].lower())]
    by_proj = defaultdict(lambda: [0.0, 0.0])
    for a in A:
        r = rs_by.get(a["resource_id"]) or {}
        by_proj[a["project_id"]][0 if M.is_union(r.get("trade")) else 1] += a["hours"]   # split by PERSON role, like the dashboard
    proj_rows = [(ps_by[pid]["name"], ps_by[pid]["division"], u, n, u + n) for pid, (u, n) in sorted(by_proj.items(), key=lambda kv: ps_by[kv[0]]["name"]) if pid in ps_by]
    return {"days": days, "labels": [_day_label(d) for d in days], "projects_row": projects_row, "rows": rows, "by_res": by_res, "by_proj": proj_rows}


def week_csv(state, monday):
    t = week_tables(state, monday)
    buf = io.StringIO()
    w = csv.writer(buf)
    w.writerow(["Pace Resource Scheduler — week of %s" % M.week_label(monday)])
    w.writerow(["Resource"] + t["labels"])
    w.writerow(["PROJECTS"] + t["projects_row"])
    for r, cells in t["rows"]:
        w.writerow([r["name"]] + cells)
    w.writerow([])
    w.writerow(["Hours by resource"])
    w.writerow(["Resource", "Role", "Total Hours"])
    for name, role, h in t["by_res"]:
        w.writerow([name, role, M.h(h)])
    w.writerow([])
    w.writerow(["Hours by project"])
    w.writerow(["Project", "Division", "Union Hours", "Non-Union Hours", "Total"])
    for name, div, u, n, tot in t["by_proj"]:
        w.writerow([name, div, M.h(u), M.h(n), M.h(tot)])
    resp = HttpResponse("﻿" + buf.getvalue(), content_type="text/csv; charset=utf-8")
    resp["Content-Disposition"] = 'attachment; filename="pace-week-%s.csv"' % iso(monday)
    return resp


def week_xlsx(state, monday):
    import openpyxl
    from openpyxl.styles import Alignment, Font
    t = week_tables(state, monday)
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Week"
    ws.append(["Pace Resource Scheduler — week of %s" % M.week_label(monday)])
    ws["A1"].font = Font(bold=True, size=13)
    ws.append(["Resource", "Role"] + t["labels"] + ["Week hours"])
    for c in ws[2]:
        c.font = Font(bold=True)
    ws.append(["PROJECTS", ""] + t["projects_row"] + [""])
    for r, cells in t["rows"]:
        ws.append([r["name"], r["trade"]] + cells + [sum(h for n, _, h in t["by_res"] if n == r["name"])])
    for row in ws.iter_rows(min_row=3):
        for c in row:
            c.alignment = Alignment(wrap_text=True, vertical="top")
    ws.column_dimensions["A"].width = 26
    for col in "CDEFGHI":
        ws.column_dimensions[col].width = 30
    ws.freeze_panes = "C3"
    w2 = wb.create_sheet("Hours by resource")
    w2.append(["Resource", "Role", "Total Hours"])
    for name, role, h in t["by_res"]:
        w2.append([name, role, h])
    w3 = wb.create_sheet("Hours by project")
    w3.append(["Project", "Division", "Union Hours", "Non-Union Hours", "Total"])
    for row in t["by_proj"]:
        w3.append(list(row))
    w4 = wb.create_sheet("Assignments")
    w4.append(["Date", "Resource", "Role", "Project", "Division", "PM", "Phase", "Hours", "Note"])
    days = t["days"]
    for a in sorted((a for a in state["assignments"] if a["date"] and days[0] <= a["date"] <= days[6]), key=lambda a: (a["date"], a["project_name"], a["resource_id"])):
        r = state["resources_by_id"].get(a["resource_id"]) or {}
        p = state["projects_by_id"].get(a["project_id"]) or {}
        w4.append([iso(a["date"]), r.get("name", "?"), r.get("trade", ""), p.get("name", "?"), p.get("division", ""), p.get("pm", ""), a["phase_name"] or "", a["hours"], a["note"]])
    for sheet in (w2, w3, w4):
        for c in sheet[1]:
            c.font = Font(bold=True)
        sheet.column_dimensions["A"].width = 24
        sheet.column_dimensions["B"].width = 24
        sheet.column_dimensions["D"].width = 44
    out = io.BytesIO()
    wb.save(out)
    resp = HttpResponse(out.getvalue(), content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
    resp["Content-Disposition"] = 'attachment; filename="pace-week-%s.xlsx"' % iso(monday)
    return resp


def resources_html(state, monday, resource_ids=None):
    """Standalone per-resource weekly schedule (print stylesheet): Day / Project (PM) / Phase / Hours, 'PTO / Vacation'
    and '—' rows."""
    days = M.week_days(monday)
    wanted = set(int(x) for x in resource_ids) if resource_ids else None
    people = []
    for r in sorted(state["resources"], key=lambda x: x["sort_name"].lower()):
        if wanted is not None and r["id"] not in wanted:
            continue
        rows, total = [], 0.0
        for d in days:
            items = [a for a in state["assignments"] if a["resource_id"] == r["id"] and a["date"] == d]
            if M.on_pto(r, d):
                rows.append({"day": _day_label(d), "project": "PTO / Vacation", "phase": "", "hours": "", "pto": True})
            elif not items:
                rows.append({"day": _day_label(d), "project": "—", "phase": "", "hours": "", "empty": True})
            else:
                for a in sorted(items, key=lambda a: a["project_name"]):
                    p = state["projects_by_id"].get(a["project_id"]) or {}
                    rows.append({"day": _day_label(d), "project": a["project_name"], "pm": p.get("pm", ""), "phase": a["phase_name"] or "", "hours": M.h(a["hours"]),
                                 "access": "%s–%s" % (M.fmt_time(p.get("access_start")), M.fmt_time(p.get("access_end")))})
                    total += a["hours"]
        people.append({"name": r["name"], "trade": r["trade"], "total": M.h(total), "rows": rows})
    html = render_to_string("scheduling/resources_print.html", {"people": people, "week": M.week_label(monday), "monday": iso(monday), "days": days})
    resp = HttpResponse(html)
    resp["Content-Disposition"] = 'inline; filename="pace-resources-%s.html"' % iso(monday)
    return resp
