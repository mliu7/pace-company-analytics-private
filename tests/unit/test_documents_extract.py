"""Text extraction (apps/documents/extract.py) on the binary fixtures under tests/fixtures/documents/."""

import os
import sys
import unittest
from pathlib import Path

sys.path.insert(0, os.path.abspath(os.path.join(os.path.dirname(__file__), "..", "..")))
os.environ.setdefault("DJANGO_SETTINGS_MODULE", "config.settings")

from apps.documents import extract as X  # noqa: E402

FIX = Path(__file__).resolve().parents[1] / "fixtures" / "documents"


class Extract(unittest.TestCase):
    def test_pdf(self):
        r = X.extract((FIX / "sample.pdf").read_bytes(), "pdf", "sample.pdf")
        self.assertTrue(r["ok"], r["error"])
        self.assertIn(r["extractor"], ("pdfplumber", "pypdf"))
        self.assertEqual(r["pages"], 1)
        self.assertIn("Timeline Theatre", r["text"])
        self.assertEqual(r["page_offsets"], [0])

    def test_docx_paragraphs_and_tables(self):
        r = X.extract((FIX / "sample.docx").read_bytes(), "docx")
        self.assertTrue(r["ok"])
        self.assertIn("Pace Systems", r["text"])
        self.assertIn("98 inch display | 1", r["text"])

    def test_xlsx_sheets_become_pages(self):
        r = X.extract((FIX / "sample.xlsx").read_bytes(), "xlsx")
        self.assertTrue(r["ok"])
        self.assertEqual(r["pages"], 2)
        self.assertIn("[sheet] Equipment", r["text"])
        self.assertIn("LG-98UM3", r["text"])
        self.assertIn("Install\t24", r["text"])

    def test_txt(self):
        r = X.extract((FIX / "sample.txt").read_bytes(), "txt")
        self.assertTrue(r["ok"]); self.assertIn("Meeting notes", r["text"])
        self.assertEqual(X.extract("caf\xe9".encode("cp1252"), "txt")["text"], "caf\xe9")

    def test_unsupported_and_garbage_never_raise(self):
        self.assertFalse(X.extract(b"x", "zip")["ok"])
        r = X.extract(b"not really a pdf", "pdf")
        self.assertFalse(r["ok"]); self.assertTrue(r["error"])
        r = X.extract(b"not a docx", "docx")
        self.assertFalse(r["ok"])

    def test_truncation(self):
        old = X.MAX_CHARS
        X.MAX_CHARS = 50
        try:
            r = X.extract(("word " * 100).encode(), "txt")
            self.assertTrue(r["truncated"]); self.assertLessEqual(len(r["text"]), 50)
        finally:
            X.MAX_CHARS = old

    def test_supported(self):
        self.assertTrue(X.supported("PDF")); self.assertTrue(X.supported("msg")); self.assertFalse(X.supported("dwg"))
        self.assertFalse(X.OCR_AVAILABLE, "pytesseract must not be installed by PCA")


if __name__ == "__main__":
    unittest.main()


class NewFormats(unittest.TestCase):
    def test_pptx_slides_are_pages(self):
        import io
        from pptx import Presentation
        from apps.documents import extract as E
        prs = Presentation()
        s = prs.slides.add_slide(prs.slide_layouts[1]); s.shapes.title.text = "Scope"; s.placeholders[1].text_frame.text = "Install 7 WAPs"
        prs.slides.add_slide(prs.slide_layouts[1]).shapes.title.text = "Price"
        b = io.BytesIO(); prs.save(b)
        r = E.extract(b.getvalue(), "pptx")
        self.assertTrue(r["ok"]); self.assertEqual(r["pages"], 2); self.assertIn("Install 7 WAPs", r["text"]); self.assertIn("[slide 2]", r["text"])

    def test_html_wearing_xls(self):
        from apps.documents import extract as E
        r = E.extract(b"<html><body><table><tr><td>Invoice</td><td>123</td></tr></table></body></html>", "xls")
        self.assertTrue(r["ok"]); self.assertEqual(r["extractor"], "xls-html"); self.assertIn("Invoice", r["text"])

    def test_doc_never_raises(self):
        from apps.documents import extract as E
        r = E.extract(b"\x00garbage", "doc")
        self.assertIn(r["ok"], (True, False)); self.assertTrue(E.supported("doc") and E.supported("pptx") and E.supported("xls"))
