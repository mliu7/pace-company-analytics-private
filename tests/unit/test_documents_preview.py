"""In-app previews (apps/documents/preview.py): sanitising, Excel / CSV tables, Word through macOS textutil where present."""

import os
import shutil
import sys
import tempfile
import unittest
from unittest.mock import patch
from pathlib import Path

sys.path.insert(0, os.path.abspath(os.path.join(os.path.dirname(__file__), "..", "..")))
os.environ.setdefault("DJANGO_SETTINGS_MODULE", "config.settings")

from apps.documents import preview as P  # noqa: E402

FIX = Path(__file__).resolve().parents[1] / "fixtures" / "share_root" / "AV" / "25-4476 Timeline Theatre"


class Sanitize(unittest.TestCase):
    def test_scripts_handlers_and_external_resources_go(self):
        dirty = ('<style>p{color:red}</style><p onclick="x()">hi</p><script>alert(1)</script><a href="javascript:evil()">l</a>'
                 '<iframe src="http://x"></iframe><link rel="stylesheet" href="http://x/a.css"><img src="data:image/png;base64,AAAA">')
        clean = P.sanitize(dirty)
        self.assertIn("<style>p{color:red}</style>", clean)
        self.assertIn("<p>hi</p>", clean)
        self.assertNotIn("script", clean)
        self.assertNotIn("iframe", clean)
        self.assertNotIn("<link", clean)
        self.assertNotIn("javascript:", clean)
        self.assertIn('<img src="data:image/png;base64,AAAA">', clean)

    @patch.object(P, "office_binary", return_value=None)
    def test_kind(self, _office):
        self.assertEqual(P.kind("xlsx"), "html")
        self.assertEqual(P.kind("csv"), "html")
        self.assertIsNone(P.kind("pdf"))
        self.assertIsNone(P.kind("zip"))
        self.assertEqual(P.kind("DOCX"), "html")
        if shutil.which("textutil"):
            self.assertEqual(P.kind("doc"), "html")
            self.assertEqual(P.kind("DOCX"), "html")


class Tables(unittest.TestCase):
    def test_xlsx_to_html(self):
        out = P.sheet_to_html(str(FIX / "Equipment list.xlsx"))
        self.assertIn('<table class="sheet">', out)
        self.assertIn("<th>A</th>", out)

    def test_csv_to_html_and_truncation(self):
        with tempfile.TemporaryDirectory() as td:
            p = Path(td) / "x.csv"
            p.write_text("a,b\n" + "\n".join("%d,%d" % (i, i * 2) for i in range(P.MAX_SHEET_ROWS + 5)))
            out = P.csv_to_html(str(p))
            self.assertIn("<td>a</td><td>b</td>", out)
            self.assertIn("truncated at %d rows" % P.MAX_SHEET_ROWS, out)

    def test_full_document_wraps_a_note(self):
        doc = P.to_html(str(FIX / "Equipment list.xlsx"), "xlsx")
        self.assertTrue(doc.startswith("<!doctype html>"))
        self.assertIn("Cell values as last saved", doc)
        self.assertIsNone(P.to_html(str(FIX / "Proposal 254476 Rev 1.pdf"), "pdf"))


@unittest.skipUnless(shutil.which("textutil"), "macOS textutil not available")
class Word(unittest.TestCase):
    def test_docx_to_html(self):
        out = P.word_to_html(str(FIX / "25-4476 Proposal.docx"), "docx")
        self.assertIsNotNone(out)
        self.assertNotIn("<script", out.lower())
        self.assertTrue(len(out) > 20)

    def test_garbage_does_not_raise(self):
        """textutil reads unknown bytes as plain text rather than failing — either a string or None, never an exception."""
        with tempfile.TemporaryDirectory() as td:
            p = Path(td) / "bad.doc"
            p.write_bytes(b"\x00\x01not a word file")
            out = P.word_to_html(str(p), "doc")
            self.assertTrue(out is None or isinstance(out, str))


if __name__ == "__main__":
    unittest.main()


class Slides(unittest.TestCase):
    def test_pptx_to_html(self):
        from pptx import Presentation
        with tempfile.TemporaryDirectory() as td:
            prs = Presentation()
            s = prs.slides.add_slide(prs.slide_layouts[1])
            s.shapes.title.text = "Scope & Approach"
            s.placeholders[1].text_frame.text = "Install 7 WAPs <1st FL>"
            p = Path(td) / "deck.pptx"
            prs.save(str(p))
            out = P.pptx_to_html(str(p))
            self.assertIn("Slide 1 · Scope &amp; Approach", out)
            self.assertIn("Install 7 WAPs &lt;1st FL&gt;", out)
            with patch.object(P, "office_binary", return_value=None):
                self.assertEqual(P.kind("pptx"), "html")
            doc = P.to_html(str(p), "pptx")
            self.assertIn("Slide text only", doc)

    def test_quicklook_is_apple_native_only(self):
        self.assertNotIn("pptx", P.QUICKLOOK_EXTS)
        self.assertNotIn("xls", P.QUICKLOOK_EXTS)


class Fallbacks(unittest.TestCase):
    @patch("shutil.which", return_value=None)
    def test_docx_on_linux(self, _which):
        out = P.word_to_html(str(FIX / "25-4476 Proposal.docx"), "docx")
        self.assertIn("<p>", out)
        self.assertGreater(len(out), 20)
        with tempfile.TemporaryDirectory() as td:
            bad = Path(td) / "bad.docx"
            bad.write_bytes(b"not a zip")
            self.assertIsNone(P.word_to_html(str(bad), "docx"))

    def test_tiff_is_converted_to_browser_image(self):
        from PIL import Image
        with tempfile.TemporaryDirectory() as td:
            source, output = Path(td) / "scan.tiff", Path(td) / "page.png"
            Image.new("RGB", (2400, 1200), "red").save(source)
            self.assertEqual(P.kind("tiff"), "image")
            self.assertTrue(P.to_png(source, "tiff", output))
            with Image.open(output) as image:
                self.assertEqual(image.format, "PNG")
                self.assertEqual(image.size, (1600, 800))

    @patch.object(P, "office_binary", return_value="/usr/bin/libreoffice")
    def test_office_formats_and_converter_failure(self, _office):
        for ext in P.OFFICE_EXTS:
            self.assertEqual(P.kind(ext.upper()), "pdf")
        with tempfile.TemporaryDirectory() as td, patch.object(P, "_run", return_value=P._Failed()):
            source, output = Path(td) / "file.doc", Path(td) / "preview.pdf"
            source.write_bytes(b"invalid")
            self.assertFalse(P.to_pdf(source, "doc", output))
            self.assertFalse(output.exists())

    @unittest.skipUnless(P.office_binary(), "LibreOffice not configured")
    def test_office_renders_word_excel_and_slides(self):
        from pptx import Presentation
        with tempfile.TemporaryDirectory() as td:
            deck = Path(td) / "slides.pptx"
            prs = Presentation()
            prs.slides.add_slide(prs.slide_layouts[0]).shapes.title.text = "Preview test"
            prs.save(deck)
            for source in (FIX / "25-4476 Proposal.docx", FIX / "Equipment list.xlsx", deck):
                output = Path(td) / (source.stem + ".pdf")
                self.assertTrue(P.to_pdf(source, source.suffix[1:], output), source.name)
                self.assertTrue(output.read_bytes().startswith(b"%PDF-"))

    def test_html_wearing_xls(self):
        with tempfile.TemporaryDirectory() as td:
            p = Path(td) / "export.xls"
            p.write_bytes(b"<html><head><script>x()</script></head><body><table><tr><td>Invoice</td><td>123</td></tr></table></body></html>")
            out = P.xls_to_html(str(p))
            self.assertIn('<table class="sheet">', out)
            self.assertIn("<td>Invoice</td>", out)
            self.assertNotIn("script", out)

    def test_raw_xlsx_reader_matches_openpyxl(self):
        raw = P.xlsx_raw_to_html(str(FIX / "Equipment list.xlsx"))
        via = P.sheet_to_html(str(FIX / "Equipment list.xlsx"))
        import re
        cells = lambda s: re.findall(r"<td[^>]*>([^<]*)</td>", s)
        self.assertEqual([c for c in cells(raw) if c], [c for c in cells(via) if c])
