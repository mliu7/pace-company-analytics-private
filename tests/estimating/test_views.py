"""Estimating workbench — Django tests (SharePoint spec §8, §11): the write gate (estimating.write), versioned saves
with the stale-form check, quick add, the bid link and estimates_for_bid, exports, rates, catalog search, import."""

import io
import json
from decimal import Decimal

from django.contrib.auth.models import User
from django.test import Client, TestCase

from apps.access.models import Account, RoleAssignment
from apps.bids.models import Bid
from apps.estimating import api, loaders, services
from apps.estimating.models import CatalogItem, CatalogSource, Estimate, EstimateVersion, LaborRate

SAVE = "/estimating/estimates/save/"


def mk(email, name, role=None, superadmin=False):
    user = User.objects.create_user(email, email=email)
    user.set_unusable_password()
    user.save()
    a = Account.objects.create(email=email, display_name=name, user=user, is_superadmin=superadmin)
    if role:
        RoleAssignment.objects.create(account=a, role=role, division_codes=None)
    return a


def client_for(acct):
    c = Client()
    c.force_login(acct.user)
    return c


class EstimatingTests(TestCase):
    @classmethod
    def setUpTestData(cls):
        cls.estimator = mk("t-est@t.local", "T Estimator", "estimator")          # estimating.view + estimating.write
        cls.finance = mk("t-fin@t.local", "T Finance", "finance")                 # estimating.view only
        cls.sales = mk("t-sales@t.local", "T Sales", "sales")                     # neither
        cls.sup = mk("t-sup@t.local", "T Owner", superadmin=True)
        src = CatalogSource.objects.create(name="Alfatron_2025-08-18.xlsx", vendor="Alfatron", kind="bootstrap", date_key=20250818)
        cls.item = CatalogItem.objects.create(source=src, manufacturer="Alfatron", manufacturer_raw="Alfatron", part="ALF-SC61E", part_norm="ALFSC61E",
                                              description="Switcher with 6 inputs", cost=Decimal("396"), msrp=Decimal("660"), date_key=20250818, confidence=90)
        cls.item2 = CatalogItem.objects.create(source=src, manufacturer="West Penn Wire", manufacturer_raw="WEST PENN", part="CAT6-BULK", part_norm="CAT6BULK",
                                               description="Cat6 bulk cable 1000 ft", cost=Decimal("120"), msrp=None, date_key=20250818, confidence=60)
        cls.bid = Bid.objects.create(sp_item_id="900", project_name="Test Lobby AV", client_name="Acme", stage="submitted", budget=Decimal("9000"), value=Decimal("12000"))

    # ---- access
    def test_view_gate(self):
        for url in ("/estimating/", "/estimating/estimates/", "/estimating/rates/", "/estimating/sources/", "/estimating/search/?q=alf"):
            self.assertEqual(client_for(self.finance).get(url).status_code, 200, url)
            self.assertEqual(client_for(self.sales).get(url).status_code, 403, url)
            self.assertEqual(Client().get(url).status_code, 302, url)
        # the import page writes (registry: estimating.view + estimating.write) — view-only roles are refused
        self.assertEqual(client_for(self.finance).get("/estimating/import/").status_code, 403)
        self.assertEqual(client_for(self.estimator).get("/estimating/import/").status_code, 200)

    def test_write_gate_on_save_rates_import(self):
        c_view, c_write = client_for(self.finance), client_for(self.estimator)
        self.assertEqual(c_view.post(SAVE, {"action": "create", "title": "x"}).status_code, 403)
        self.assertEqual(c_view.post("/estimating/rates/save/", {"cost_engineering": "80"}).status_code, 403)
        self.assertEqual(c_view.post("/estimating/import/", {"action": "remove_source", "source_id": 1}).status_code, 403)
        r = c_write.post(SAVE, {"action": "create", "title": "Mine"})
        self.assertEqual(r.status_code, 302)
        self.assertEqual(Estimate.objects.get(title="Mine").owner, self.estimator)
        # a view-only user sees the builder read-only (no Save button, inputs read-only)
        est = Estimate.objects.get(title="Mine")
        html = c_view.get("/estimating/estimates/%d/" % est.id).content.decode()
        self.assertNotIn('id="btn-save"', html)
        self.assertIn("readonly", html)
        self.assertIn('id="btn-save"', c_write.get("/estimating/estimates/%d/" % est.id).content.decode())

    # ---- versioned save
    def _create(self, c, title="T"):
        r = c.post(SAVE, data=json.dumps({"action": "create", "title": title, "client_name": "Acme"}), content_type="application/json")
        self.assertEqual(r.status_code, 200, r.content)
        return Estimate.objects.get(pk=r.json()["id"])

    def _payload(self, est):
        return {"action": "save", "id": est.id, "version_no": est.version_no, "title": est.title, "client_name": "Acme", "notes": "", "status": "draft",
                "rooms": [{"name": "Conference A", "lines": [
                    {"catalog_item_id": self.item.id, "manufacturer": "Alfatron", "part": "ALF-SC61E", "description": "Switcher", "qty": 3, "cost": 396, "markup": 1.4,
                     "hours": {"field_labor": 10, "engineering": 2.5}},
                    {"manufacturer": "West Penn", "part": "CAT6-BULK", "description": "Cat6 bulk cable", "qty": 2, "cost": 120, "markup": 1.2, "hours": {"union_pull": 90}},
                    {"kind": "note", "description": "Owner supplies mounts"}]},
                    {"name": "Lobby", "lines": [{"part": "ALF-SC61E", "description": "Switcher", "qty": 1, "cost": 396}]}]}

    def test_save_computes_versions_and_refuses_stale(self):
        c = client_for(self.estimator)
        est = self._create(c)
        self.assertEqual(est.version_no, 0)
        r = c.post(SAVE, data=json.dumps(self._payload(est)), content_type="application/json")
        self.assertEqual(r.status_code, 200, r.content)
        p = r.json()["payload"]
        self.assertEqual(p["version_no"], 1)
        self.assertEqual([x["name"] for x in p["rooms"]], ["Conference A", "Lobby"])
        l0 = p["rooms"][0]["lines"][0]
        self.assertEqual(l0["sell"], 554.4)                                  # 396 × 1.4
        self.assertEqual(l0["markup"], 1.4)
        self.assertEqual(l0["totals"]["equipment_cost_ext"], 1188)         # × qty
        self.assertEqual(l0["totals"]["labor_sell"], 10 * 125 + 2.5 * 135)  # hours × default sell rate, NOT × qty
        misc = p["rooms"][0]["lines"][1]
        self.assertTrue(misc["misc"])
        self.assertEqual(misc["totals"]["labor_cost"], 90 * 92)
        self.assertEqual(p["rooms"][0]["lines"][2]["kind"], "note")
        self.assertEqual(p["rooms"][1]["lines"][0]["sell"], 500.94)         # default 1.265 when only cost came
        codes = {w["code"] for w in p["warnings"]}
        self.assertEqual(codes, {"peer_review", "consumables_markup"})        # 102.5 h > 80; CAT6 at 1.2 < 1.5
        est.refresh_from_db()
        self.assertEqual(est.line_count, 4)
        self.assertEqual(est.room_count, 2)
        self.assertEqual(float(est.total_sell), p["grand"]["total_sell"])
        self.assertEqual(EstimateVersion.objects.filter(estimate=est).count(), 1)
        self.assertEqual(est.versions.first().snapshot["rooms"][0]["lines"][0]["part"], "ALF-SC61E")
        # stale: same payload (version_no 0) again → 409 with who / when
        r2 = c.post(SAVE, data=json.dumps(self._payload(Estimate(id=est.id, version_no=0, title="T"))), content_type="application/json")
        self.assertEqual(r2.status_code, 409)
        self.assertIn("changed by T Estimator", r2.json()["error"])
        # fresh version_no saves again → v2; restore v1 → v3
        pl = self._payload(est)
        pl["rooms"] = pl["rooms"][:1]
        r3 = c.post(SAVE, data=json.dumps(pl), content_type="application/json")
        self.assertEqual(r3.json()["payload"]["version_no"], 2)
        r4 = c.post(SAVE, data=json.dumps({"action": "restore_version", "id": est.id, "version_no": 1}), content_type="application/json")
        self.assertEqual(len(r4.json()["payload"]["rooms"]), 2)
        self.assertEqual(r4.json()["payload"]["version_no"], 3)

    def test_quick_add_duplicate_delete_and_bid_link(self):
        c = client_for(self.estimator)
        est = self._create(c)
        r = c.post(SAVE, data=json.dumps({"action": "add_line", "id": est.id, "item_id": self.item.id, "qty": 2}), content_type="application/json")
        self.assertEqual(r.status_code, 200, r.content)
        self.assertEqual(r.json()["estimate"]["total_sell"], 1001.88)        # 2 × 500.94
        est.refresh_from_db()
        self.assertEqual(est.version_no, 1)
        self.assertEqual(est.rooms.first().lines.first().catalog_item_id, self.item.id)
        # attach the bid: client / customer follow, summaries expose budget vs sell
        r = c.post(SAVE, data=json.dumps({"action": "attach_bid", "id": est.id, "bid_id": self.bid.id}), content_type="application/json")
        self.assertEqual(r.json()["payload"]["bid"]["value"], 12000.0)
        self.assertEqual(r.json()["payload"]["bid"]["budget"], 9000.0)
        summaries = api.estimates_for_bid(self.bid.id)
        self.assertEqual(len(summaries), 1)
        self.assertEqual(summaries[0]["total_sell"], Decimal("1001.88"))
        self.assertEqual(summaries[0]["url"], "/estimating/estimates/%d/" % est.id)
        html = c.get("/estimating/estimates/%d/" % est.id).content.decode()
        self.assertIn("Test Lobby AV", html)
        # approval bridge: live when Phase D's API exists (a Pending ApprovalRequest linked to the estimate + bid), guarded otherwise
        est.refresh_from_db()                                                # attach_bid happened through the view
        ok, msg = api.raise_approval(est, self.estimator, note="please rush", kind="Labor", needed_by="2026-10-01")
        if api.approval_available():
            self.assertTrue(ok, msg)
            from apps.planning.models import ApprovalRequest
            req = ApprovalRequest.objects.get(estimate_id=est.id)
            self.assertEqual((req.kind, req.bid_id, req.status), ("Labor", self.bid.id, ApprovalRequest.Status.PENDING))
            self.assertIn("please rush", req.notes)
            self.assertIn("sell $1,001.88", req.notes)
            self.assertEqual(str(req.needed_by), "2026-10-01")
            est.refresh_from_db()
            self.assertEqual(est.approval_ref, str(req.pk))
            r = c.post(SAVE, data=json.dumps({"action": "approval", "id": est.id, "kind": "BOM"}), content_type="application/json")
            self.assertTrue(r.json()["ok"], r.content)
            self.assertEqual(ApprovalRequest.objects.filter(estimate_id=est.id).count(), 2)
        else:
            self.assertIn("Production phase", msg)
        # duplicate + delete
        r = c.post(SAVE, data=json.dumps({"action": "duplicate", "id": est.id}), content_type="application/json")
        copy = Estimate.objects.get(pk=r.json()["id"])
        self.assertEqual(copy.title, "T Copy")
        self.assertEqual(copy.line_count, 1)
        self.assertEqual(copy.bid_id, self.bid.id)
        r = c.post(SAVE, data=json.dumps({"action": "delete", "id": copy.id}), content_type="application/json")
        self.assertTrue(r.json()["ok"])
        self.assertFalse(Estimate.objects.filter(pk=copy.id).exists())

    def test_exports_and_search(self):
        c = client_for(self.estimator)
        est = self._create(c)
        c.post(SAVE, data=json.dumps(self._payload(est)), content_type="application/json")
        x = c.get("/estimating/estimates/%d/export/?fmt=xlsx" % est.id)
        self.assertEqual(x.status_code, 200)
        from openpyxl import load_workbook
        wb = load_workbook(io.BytesIO(x.content))
        self.assertEqual(wb.sheetnames, ["Summary", "Conference A", "Lobby"])
        self.assertEqual(wb["Summary"]["A7"].value, "Room")
        csv_ = c.get("/estimating/estimates/%d/export/?fmt=csv" % est.id).content.decode()
        self.assertIn("Field Labor Sell Extended", csv_)
        self.assertEqual(csv_.count("\n"), 1 + 4 + 2)                        # header + 4 lines + 2 room totals
        txt = c.get("/estimating/estimates/%d/export/?fmt=txt" % est.id).content.decode()
        self.assertIn("TOTAL: Cost", txt)
        s = c.get("/estimating/search/?q=alf-sc61e").json()
        self.assertEqual(s["total"], 1)
        self.assertEqual(s["rows"][0]["score_label"], "Exact")
        self.assertEqual(s["rows"][0]["margin"], 40.0)
        s2 = c.get("/estimating/search/?q=cable&has_cost=1&sort=cost_desc").json()
        self.assertEqual(s2["rows"][0]["part"], "CAT6-BULK")
        self.assertEqual(c.get("/estimating/search/?q=").json()["total"], 0)
        self.assertEqual(c.get("/estimating/search/?brand=Alfatron").json()["total"], 1)
        self.assertEqual(c.get("/estimating/search/?kind=bids&q=lobby").json()["rows"][0]["id"], self.bid.id)
        self.assertEqual(c.get("/estimating/item/%d/" % self.item.id).status_code, 200)

    def test_rates_effective_dated(self):
        c = client_for(self.estimator)
        self.assertEqual(services.current_rates()["engineering"]["sell"], Decimal("135.00"))
        r = c.post("/estimating/rates/save/", {"effective_from": "2026-01-02", "cost_engineering": "80", "sell_engineering": "155"})
        self.assertEqual(r.status_code, 302)
        self.assertEqual(services.current_rates()["engineering"]["sell"], Decimal("155.00"))
        self.assertEqual(services.current_rates()["union_pull"]["sell"], Decimal("125.00"))
        self.assertEqual(LaborRate.objects.filter(rate_id="engineering").count(), 2)
        c.post("/estimating/rates/save/", {"effective_from": "2026-01-03", "action": "reset"})
        self.assertEqual(services.current_rates()["engineering"]["sell"], Decimal("135.00"))
        self.assertEqual(c.get("/estimating/rates/").status_code, 200)

    def test_catalog_import_modes_and_hygiene(self):
        matrix = [["Manufacturer", "Part Number", "Description", "Dealer Cost", "MSRP / List Price", "MAP Price", "Effective Date"],
                  ["Alfatron", "ALF-SC61E", "Switcher new price", "380", "660", "", "2026-01-15"],
                  ["Alfatron", "ALF-NEW1", "New product", "299", "499", "", "2026-01-15"],
                  ["Alfatron", "ALF-OLD", "Ancient", "10", "", "", "2019-01-01"]]
        rows = loaders_rows(matrix, "Alfatron_2026-01-15.xlsx")
        dry = loaders.apply_catalog_rows(rows, "Alfatron_2026-01-15.xlsx", mode="upsert", dry_run=True)
        self.assertEqual((dry["counts"]["added"], dry["counts"]["updated"], dry["counts"]["stale"]), (1, 1, 1))
        self.assertEqual(CatalogItem.objects.count(), 2)
        res = loaders.apply_catalog_rows(rows, "Alfatron_2026-01-15.xlsx", mode="add-only", dry_run=False)
        self.assertEqual((res["counts"]["added"], res["counts"]["updated"]), (1, 0))
        self.assertTrue(CatalogItem.objects.filter(part_norm="ALFNEW1", archived=False).exists())
        res = loaders.apply_catalog_rows(rows, "Alfatron_2026-01-15b.xlsx", mode="prices-only", dry_run=False)
        self.assertEqual((res["counts"]["added"], res["counts"]["updated"]), (0, 1))
        cur = CatalogItem.objects.get(part_norm="ALFSC61E", archived=False)
        self.assertEqual(cur.cost, Decimal("380"))                             # the newer file's row is current
        self.assertEqual(cur.date_key, 20260115)
        old = CatalogItem.objects.get(pk=self.item.pk)
        self.assertTrue(old.archived)
        self.assertEqual(old.archive_reason, "duplicate")
        self.assertEqual(old.superseded_by_id, cur.id)
        self.assertEqual(res["hygiene"]["current"], 3)
        # removing the upload brings the old row back
        loaders.remove_source(res["source_id"])
        old.refresh_from_db()
        self.assertFalse(old.archived)
        self.assertFalse(CatalogItem.objects.filter(part_norm="ALFSC61E", cost=Decimal("380")).exists())

    def test_estimate_import_prices_from_catalog(self):
        sheets = [("Lobby", [["", "Lobby and Total Counts"], ["Item", "Manufacturer", "Model #", "Description", "Qty", "Cost"],
                             ["Video", "Alfatron", "ALF-SC61E", "Switcher", "2", ""], ["Cable", "", "CAT6-BULK", "bulk cable", "1", "100"],
                             ["", "", "", "Owner note", "0", ""], ["", "Nobody", "UNKNOWN-1", "mystery", "1", "50"]])]
        res = services.import_rooms(sheets, "Imported estimate: bom.xlsx")
        self.assertEqual(res["strategy"], "sheet_per_room")
        self.assertEqual(res["rooms"][0]["name"], "Lobby")
        lines = res["rooms"][0]["lines"]
        self.assertEqual((res["matched"], res["unmatched"], res["notes"]), (2, 1, 1))
        self.assertEqual(lines[0]["cost"], Decimal("396"))                     # catalog cost
        self.assertEqual(lines[0]["catalog_item_id"], self.item.id)
        self.assertEqual(lines[1]["cost"], Decimal("100"))                     # file cost overrides the catalog's 120
        self.assertEqual(lines[3]["catalog_item_id"], None)
        c = client_for(self.estimator)
        est = self._create(c, "Imported")
        est = services.save_estimate(est, {"rooms": res["rooms"]}, self.estimator, note="import")
        self.assertEqual(est.line_count, 4)
        self.assertEqual(est.rooms.first().lines.filter(kind="note").count(), 1)


def loaders_rows(matrix, name):
    from apps.estimating import rules
    return rules.catalog_rows(matrix, name, "Sheet1")
